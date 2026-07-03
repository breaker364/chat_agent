"""
commands/bitable.py - Bitable read commands (mirrors lib/commands/bitable.js)
"""

import base64
import json
import os
import re
import sys
import zlib
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone

from ..config import BITABLE_HOST, DOC_HOST
from ..http_utils import http_get, http_post_with_cookies


# ---------------------------------------------------------------------------
# resolve_bitable_token
# ---------------------------------------------------------------------------

def resolve_bitable_token(input_str: str) -> dict:
    """Parse bitable URL/token. Returns {token, is_wiki, table_id, view_id}."""
    result = {'token': input_str, 'is_wiki': False, 'table_id': None, 'view_id': None}
    if '/' not in input_str:
        return result
    # extract query params
    from urllib.parse import urlparse, parse_qs
    parsed = urlparse(input_str)
    qs = parse_qs(parsed.query)
    result['table_id'] = qs.get('table', [None])[0]
    result['view_id'] = qs.get('view', [None])[0]
    m = re.search(r'/base/([A-Za-z0-9]+)', input_str)
    if m:
        result['token'] = m.group(1)
        return result
    m2 = re.search(r'/(wiki|bitable)/([A-Za-z0-9]+)', input_str)
    if m2:
        result['is_wiki'] = m2.group(1) == 'wiki'
        result['token'] = m2.group(2)
    return result


def resolve_wiki_to_bitable(cookies, wiki_token: str) -> str:
    """Resolve wiki node token to bitable obj_token via wiki API."""
    node_res = http_get(cookies, DOC_HOST, f'/space/api/wiki/v2/tree/get_node/?wiki_token={wiki_token}&expand_shortcut=true')
    node_data = (node_res.get('data') or {}).get('data') or {}
    if (node_res.get('data') or {}).get('code') == 0 and node_data.get('obj_token'):
        print(f"[bitable] Wiki resolved: obj_token={node_data['obj_token']}", file=sys.stderr)
        return node_data['obj_token']
    raise RuntimeError(f'Wiki resolution failed for token: {wiki_token}')


# ---------------------------------------------------------------------------
# fetch_bitable_schema
# ---------------------------------------------------------------------------

def fetch_bitable_schema(cookies, token: str, table_id: str = None) -> dict:
    if not table_id:
        url = (f'/space/api/v1/bitable/{token}/clientvars?tableID=&viewID='
               f'&recordLimit=0&ondemandLimit=0&needBase=true&viewLazyLoad=true'
               f'&ondemandVer=2&openType=1&noMissCS=true&optimizationFlag=1&removeFmlExtra=true')
        res = http_get(cookies, BITABLE_HOST, url)
        payload = res.get('data')
        if not isinstance(payload, dict) or payload.get('code') != 0:
            raise RuntimeError(f"Failed to fetch base info: payload={payload!r}")
        base_str = payload.get('data', {}).get('base')
        if not base_str:
            raise RuntimeError('No base field in clientvars response')
        base_json = json.loads(zlib.decompress(base64.b64decode(base_str), 47))  # 47=wbits for gzip
        table_map = {}
        for tid in (base_json.get('blocks') or []):
            if not tid.startswith('tbl'):
                continue
            info = (base_json.get('blockInfos') or {}).get(tid) or {}
            table_map[tid] = {'name': info.get('name') or tid}
        return {'tableMap': table_map}
    else:
        url = f'/space/api/bitable/{token}/tablesv3/'
        body = {
            'tableIDList': [table_id],
            'tablePartitionFlagList': [0],
            'tablePartitionForNoRankFlagList': [],
            'encodingProtocol': {'compression': 1, 'serialization': 0},
        }
        res = http_post_with_cookies(cookies, BITABLE_HOST, url, body)
        payload = res.get('data')
        if not isinstance(payload, dict) or payload.get('code') != 0:
            raise RuntimeError(f"Failed to fetch table schema: payload={payload!r}")
        encoded = payload.get('data', {}).get(table_id)
        if not encoded:
            raise RuntimeError(f'No schema data for table {table_id}')
        table_data = json.loads(zlib.decompress(base64.b64decode(encoded), 47))
        return {'data': {'table': table_data}}


# ---------------------------------------------------------------------------
# fetch_bitable_records
# ---------------------------------------------------------------------------

def fetch_bitable_records(cookies, token: str, table_id: str, offset: int, limit: int) -> dict:
    url = (f'/space/api/v1/bitable/{token}/records?tableId={table_id}'
           f'&viewLazyLoad=true&offset={offset}&limit={limit}'
           f'&tableID={table_id}&removeFmlExtra=true')
    res = http_get(cookies, BITABLE_HOST, url)
    if not res.get('data') or res['data'].get('code') != 0:
        raise RuntimeError(f"Failed to fetch records: code={(res.get('data') or {}).get('code')}")
    d = res['data'].get('data') or {}
    parsed = None
    if d.get('records'):
        buf = base64.b64decode(d['records'])
        parsed = json.loads(zlib.decompress(buf, 47))
    total_num = (parsed or {}).get('tableRecordNum') or (parsed or {}).get('viewRecordNum') or 0
    record_map = (parsed or {}).get('recordMap') or {}
    record_count = len(record_map)
    return {'recordMap': record_map, 'total': total_num, 'hasMore': record_count >= limit, 'offset': offset}


def fetch_all_bitable_records(cookies, token: str, table_id: str,
                              page_size: int = 3000, max_workers: int = 8) -> dict:
    """Fetch every record in a table.

    Strategy: synchronously fetch page 1 to learn `tableRecordNum`, then
    issue all remaining pages concurrently via a thread pool. Server cap
    on `limit` is 3000 (probed 2026-05-07 against a 19515-row table; 5000
    returns code=800004006). 8-way concurrency was clean (no rate-limit).
    """
    first = fetch_bitable_records(cookies, token, table_id, 0, page_size)
    total = first.get('total') or 0
    merged = dict(first.get('recordMap') or {})
    if not merged or len(merged) >= total:
        return {'recordMap': merged, 'total': total or len(merged)}

    offsets = list(range(len(merged), total, page_size))
    with ThreadPoolExecutor(max_workers=max_workers) as ex:
        futures = [ex.submit(fetch_bitable_records, cookies, token, table_id, off, page_size)
                   for off in offsets]
        for f in as_completed(futures):
            merged.update(f.result().get('recordMap') or {})

    return {'recordMap': merged, 'total': total or len(merged)}


# ---------------------------------------------------------------------------
# build_field_map
# ---------------------------------------------------------------------------

def build_field_map(schema: dict) -> dict:
    fm = ((schema.get('data') or {}).get('table') or {}).get('fieldMap') or {}
    result = {}
    for fid, fdef in fm.items():
        option_map = {}
        for opt in ((fdef.get('property') or {}).get('options') or []):
            option_map[opt['id']] = opt['name']
        result[fid] = {'name': fdef.get('name') or fid, 'type': fdef.get('type') or 0, 'optionMap': option_map}
    return result


# ---------------------------------------------------------------------------
# extract_cell_value
# ---------------------------------------------------------------------------

def extract_cell_value(cell, field_info: dict):
    if cell is None:
        return None
    val = cell.get('value') if isinstance(cell, dict) and 'value' in cell else cell
    if val is None:
        return None
    ftype = (field_info or {}).get('type', 0)

    if ftype in (5, 20, 1001, 1002) and isinstance(val, (int, float)):
        return datetime.fromtimestamp(val / 1000).astimezone().isoformat(timespec='seconds')

    if ftype == 3 and isinstance(val, str):
        return (field_info.get('optionMap') or {}).get(val, val)

    if ftype == 4 and isinstance(val, list):
        om = field_info.get('optionMap') or {}
        return ', '.join(om.get(v, v) for v in val)

    if ftype == 11:
        parsed = val
        if isinstance(val, str):
            try:
                parsed = json.loads(val)
            except Exception:
                pass
        if isinstance(parsed, dict) and parsed.get('users'):
            return ', '.join(u.get('name') or u.get('enName') or u.get('en_name') or u.get('userId', '') for u in parsed['users'])

    if isinstance(val, list):
        parts = []
        for v in val:
            if isinstance(v, dict):
                parts.append(v.get('name') or v.get('en_name') or v.get('text') or v.get('id') or json.dumps(v))
            else:
                parts.append(str(v))
        return ', '.join(parts)

    if isinstance(val, dict):
        if val.get('text') is not None:
            return val['text']
        if val.get('link') is not None:
            return val['link']
        return json.dumps(val)

    return val


# ---------------------------------------------------------------------------
# records_to_rows
# ---------------------------------------------------------------------------

def records_to_rows(record_map: dict, field_map: dict) -> list:
    rows = []
    for rid, rec in record_map.items():
        if not isinstance(rec, dict):
            continue
        row = {'_recordId': rid}
        for fid, finfo in field_map.items():
            row[finfo['name']] = extract_cell_value(rec.get(fid), finfo)
        rows.append(row)
    return rows


# ---------------------------------------------------------------------------
# fetch_bitable_field_map_batch
# ---------------------------------------------------------------------------

def fetch_bitable_field_map_batch(cookies, token: str, table_ids: list) -> dict:
    url = f'/space/api/bitable/{token}/tablesv3/'
    body = {
        'tableIDList': table_ids,
        'tablePartitionFlagList': [0] * len(table_ids),
        'tablePartitionForNoRankFlagList': [],
        'encodingProtocol': {'compression': 1, 'serialization': 0},
    }
    res = http_post_with_cookies(cookies, BITABLE_HOST, url, body)
    payload = res.get('data')
    if not isinstance(payload, dict) or payload.get('code') != 0:
        raise RuntimeError(f"Failed to fetch table schemas: payload={payload!r}")
    result = {}
    for tid in table_ids:
        encoded = payload.get('data', {}).get(tid)
        if not encoded:
            continue
        try:
            result[tid] = json.loads(zlib.decompress(base64.b64decode(encoded), 47))
        except Exception:
            pass
    return result


# ---------------------------------------------------------------------------
# BITABLE_FIELD_TYPES + format_field_schema
# ---------------------------------------------------------------------------

ROLLUP_NAMES = {0: 'original', 1: 'count', 2: 'counta', 3: 'countall', 4: 'sum', 5: 'average', 6: 'lookup'}

# BITABLE_FIELD_TYPES + VIEW_TYPE_NAMES are owned by bitable_query.py (since
# parse_filter_arg / parse_view_config consume them) and re-exported below.
# format_field_schema still references BITABLE_FIELD_TYPES via the re-export.


def format_field_schema(fid: str, fdef: dict) -> dict:
    ui_type = fdef.get('fieldUIType') or BITABLE_FIELD_TYPES.get(fdef.get('type'), 'Unknown')
    field = {
        'fieldId': fid,
        'name': fdef.get('name'),
        'type': fdef.get('type'),
        'uiType': ui_type,
        'isPrimary': fdef.get('isPrimary', False),
    }
    desc = fdef.get('description') or {}
    if desc.get('content'):
        field['description'] = ''.join(c.get('text', '') for c in desc['content'])
    p = fdef.get('property')
    if not p:
        return field
    if p.get('options'):
        field['options'] = [o['name'] for o in p['options']]
    if p.get('formula'):
        field['formula'] = p['formula']
    if p.get('formatter') and fdef.get('type') == 2:
        field['format'] = p['formatter']
    if p.get('dateFormat'):
        field['dateFormat'] = p['dateFormat']
        if p.get('timeFormat'):
            field['timeFormat'] = p['timeFormat']
    if p.get('tableId'):
        field['linkedTable'] = p['tableId']
    if p.get('targetField'):
        field['lookupField'] = p['targetField']
    if p.get('multiple') is not None:
        field['multiple'] = p['multiple']
    if p.get('rollup') is not None:
        field['rollup'] = ROLLUP_NAMES.get(p['rollup'], p['rollup'])
    return field


# ---------------------------------------------------------------------------
# cmd_bitable_tables
# ---------------------------------------------------------------------------

def cmd_bitable_tables(cookies, token: str):
    schema = fetch_bitable_schema(cookies, token)
    table_map = schema.get('tableMap') or {}
    table_ids = list(table_map.keys())
    all_schemas = fetch_bitable_field_map_batch(cookies, token, table_ids)
    tables = []
    for tid in table_ids:
        tdef = table_map[tid]
        t_data = all_schemas.get(tid) or {}
        fm = t_data.get('fieldMap') or {}
        meta = t_data.get('meta') or {}
        fields = [{'fieldId': fid, 'name': fdef.get('name') or fid, 'type': fdef.get('type') or 0} for fid, fdef in fm.items()]
        tables.append({'tableId': tid, 'name': tdef.get('name') or tid, 'recordsNum': meta.get('recordsNum') or 0, 'fieldCount': len(fields), 'fields': fields})
    print(json.dumps({'success': True, 'token': token, 'tableCount': len(tables), 'tables': tables}, indent=2))


# ---------------------------------------------------------------------------
# cmd_bitable_schema
# ---------------------------------------------------------------------------

def cmd_bitable_schema(cookies, token: str, table_id: str = None):
    base_schema = fetch_bitable_schema(cookies, token)
    table_map = base_schema.get('tableMap') or {}

    if not table_id:
        table_ids = list(table_map.keys())
        all_schemas = fetch_bitable_field_map_batch(cookies, token, table_ids)
        tables = []
        for tid in table_ids:
            tdef = table_map[tid]
            t_data = all_schemas.get(tid)
            if not t_data:
                tables.append({'tableId': tid, 'name': tdef.get('name') or tid, 'error': 'no access'})
                continue
            meta = t_data.get('meta') or {}
            fields = [format_field_schema(fid, fdef) for fid, fdef in (t_data.get('fieldMap') or {}).items()]
            tables.append({'tableId': tid, 'name': tdef.get('name') or tid, 'recordsNum': meta.get('recordsNum') or 0, 'fields': fields})
        print(json.dumps({'success': True, 'token': token, 'tableCount': len(tables), 'tables': tables}, indent=2))
    else:
        schema = fetch_bitable_schema(cookies, token, table_id)
        fm = ((schema.get('data') or {}).get('table') or {}).get('fieldMap') or {}
        meta = ((schema.get('data') or {}).get('table') or {}).get('meta') or {}
        table_name = (table_map.get(table_id) or {}).get('name') or table_id
        fields = [format_field_schema(fid, fdef) for fid, fdef in fm.items()]
        print(json.dumps({'success': True, 'table': table_name, 'tableId': table_id, 'recordsNum': meta.get('recordsNum') or 0, 'fieldCount': len(fields), 'fields': fields}, indent=2))


# ---------------------------------------------------------------------------
# View / filter / sort / group helpers — moved to commands/bitable_query.py.
# Re-exported below for back-compat with existing call sites + tests that
# do `from lark_tools.commands.bitable import eval_condition, ...`.
# ---------------------------------------------------------------------------

from .bitable_query import (  # noqa: E402,F401
    BITABLE_FIELD_TYPES, VIEW_TYPE_NAMES,
    FIELD_TYPE_OPERATORS, OPERATOR_NAMES,
    _resolve_field, find_view, parse_view_config, _display_filter_value,
    _parse_bool, _parse_datetime_to_ms, _select_opt_ids,
    parse_filter_arg, parse_sort_arg, parse_group_arg,
    _is_null, _ci_text, _text_contains,
    eval_condition, eval_filter, apply_sort, apply_group,
)


# ---------------------------------------------------------------------------
# cmd_bitable_records
# ---------------------------------------------------------------------------

def cmd_bitable_records(cookies, token: str, table_id: str, opts: dict = None):
    opts = opts or {}
    base_schema = fetch_bitable_schema(cookies, token)
    table_map = base_schema.get('tableMap') or {}
    if not table_id:
        tids = list(table_map.keys())
        if not tids:
            print(json.dumps({'error': 'No tables found'}))
            return
        table_id = tids[0]
        print(f"[bitable] Using first table: {table_id} ({(table_map[table_id] or {}).get('name') or table_id})", file=sys.stderr)

    schema = fetch_bitable_schema(cookies, token, table_id)
    table_data = (schema.get('data') or {}).get('table') or {}
    field_map = build_field_map(schema)
    view_map = table_data.get('viewMap') or {}
    meta = table_data.get('meta') or {}
    table_name = (table_map.get(table_id) or {}).get('name') or table_id

    # Resolve --view baseline
    baseline_filter, baseline_sort, baseline_group = None, [], []
    if opts.get('view'):
        _, vdef = find_view(view_map, opts['view'])
        prop = vdef.get('property') or {}
        baseline_filter = prop.get('filterInfo')
        baseline_sort = prop.get('sortInfo') or []
        baseline_group = prop.get('group') or []

    # Parse ad-hoc CLI overrides
    cli_filters = [parse_filter_arg(f, field_map) for f in (opts.get('filters') or [])]
    cli_sorts = [parse_sort_arg(s, field_map) for s in (opts.get('sorts') or [])]
    cli_group = parse_group_arg(opts['group'], field_map) if opts.get('group') else None

    final_filter = baseline_filter
    if cli_filters:
        existing = ((baseline_filter or {}).get('conditions') or [])
        final_filter = {
            'conjunction': (baseline_filter or {}).get('conjunction') or 'and',
            'conditions': existing + cli_filters,
        }
    final_sort = cli_sorts if cli_sorts else baseline_sort
    final_group = [cli_group] if cli_group else baseline_group

    has_query = bool(final_filter or final_sort or final_group)
    needs_all = bool(opts.get('all') or has_query)

    limit = opts.get('limit', 20)
    offset = opts.get('offset', 0)

    if needs_all:
        result = fetch_all_bitable_records(cookies, token, table_id)
        has_more = False
    else:
        result = fetch_bitable_records(cookies, token, table_id, offset, limit)
        has_more = result.get('hasMore', False)

    record_map = result.get('recordMap') or {}
    items = list(record_map.items())  # [(rid, raw_record), ...]

    if final_filter:
        items = [it for it in items if eval_filter(it[1], final_filter, field_map)]
    filtered_count = len(items)

    if final_sort:
        items = apply_sort(items, final_sort, field_map)

    raw_record_map = dict(items)
    rows = records_to_rows(raw_record_map, field_map)

    out = {
        'success': True,
        'table': table_name,
        'tableId': table_id,
        'total': meta.get('recordsNum') or result.get('total') or 0,
        'filtered': filtered_count,
        'fields': [f['name'] for f in field_map.values()],
    }

    if final_group:
        gfid = final_group[0]['fieldId']
        gname = field_map.get(gfid, {}).get('name') or gfid
        out['groupBy'] = gname
        out['groups'] = apply_group(rows, gfid, field_map)
        if len(final_group) > 1:
            out['_note'] = f'group has {len(final_group)} fields; only first ({gname}) applied'
    else:
        if needs_all:
            # Apply --limit slice client-side after filter/sort (recommendation D).
            if opts.get('limit') is not None:
                rows = rows[:limit]
        else:
            out['offset'] = offset
            out['limit'] = limit
            out['hasMore'] = has_more
        out['returned'] = len(rows)
        out['records'] = rows

    print(json.dumps(out, indent=2, ensure_ascii=False))


def cmd_bitable_views(cookies, token: str, table_id: str = None):
    base_schema = fetch_bitable_schema(cookies, token)
    table_map = base_schema.get('tableMap') or {}
    if not table_id:
        tids = list(table_map.keys())
        if not tids:
            print(json.dumps({'error': 'No tables found'}))
            return
        table_id = tids[0]
        print(f"[bitable] Using first table: {table_id} ({(table_map[table_id] or {}).get('name') or table_id})", file=sys.stderr)

    schema = fetch_bitable_schema(cookies, token, table_id)
    table_data = (schema.get('data') or {}).get('table') or {}
    field_map = build_field_map(schema)
    view_map = table_data.get('viewMap') or {}
    views_order = table_data.get('views') or list(view_map.keys())

    views = []
    for i, vid in enumerate(views_order):
        vdef = view_map.get(vid)
        if not vdef:
            continue
        v = parse_view_config(vdef, field_map)
        v = {'viewId': vid, **v}
        if i == 0:
            v['isDefault'] = True
        views.append(v)

    table_name = (table_map.get(table_id) or {}).get('name') or table_id
    print(json.dumps({
        'success': True,
        'table': table_name,
        'tableId': table_id,
        'viewCount': len(views),
        'views': views,
    }, indent=2, ensure_ascii=False))


# ---------------------------------------------------------------------------
# cmd_bitable_download
# ---------------------------------------------------------------------------

def cmd_bitable_download(cookies, token: str, table_id: str, output_path: str = None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    base_schema = fetch_bitable_schema(cookies, token)
    table_map = base_schema.get('tableMap') or {}
    if not table_id:
        tids = list(table_map.keys())
        if not tids:
            print(json.dumps({'error': 'No tables found'}))
            return
        table_id = tids[0]

    schema = fetch_bitable_schema(cookies, token, table_id)
    field_map = build_field_map(schema)
    table_name = (table_map.get(table_id) or {}).get('name') or table_id
    field_names = [f['name'] for f in field_map.values()]

    all_rows, offset, PAGE_SIZE, total = [], 0, 200, None
    while True:
        print(f'[bitable] Fetching records offset={offset} limit={PAGE_SIZE}...', file=sys.stderr)
        result = fetch_bitable_records(cookies, token, table_id, offset, PAGE_SIZE)
        if total is None:
            total = result['total']
        rows = records_to_rows(result['recordMap'], field_map)
        all_rows.extend(rows)
        print(f'[bitable] Got {len(rows)} records (total so far: {len(all_rows)})', file=sys.stderr)
        if not result['hasMore'] or not rows:
            break
        offset += PAGE_SIZE

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = table_name[:31]
    ws.append(field_names)
    header_font = Font(bold=True)
    header_fill = PatternFill('solid', fgColor='E2EFDA')
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    for row in all_rows:
        vals = []
        for fn in field_names:
            v = row.get(fn)
            vals.append('' if v is None else (json.dumps(v) if isinstance(v, (dict, list)) else v))
        ws.append(vals)

    # Auto-width
    for i, col in enumerate(ws.columns):
        max_len = len(field_names[i]) if i < len(field_names) else 10
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    from ..paths import resolve_output_path
    safe_name = re.sub(r'[/\\:*?"<>|]', '_', table_name)
    output_path = resolve_output_path(output_path, f'{safe_name}.xlsx')
    wb.save(output_path)
    abs_path = os.path.abspath(output_path)
    print(json.dumps({'success': True, 'table': table_name, 'records': len(all_rows), 'total': total, 'fields': len(field_names), 'path': abs_path}, indent=2))
