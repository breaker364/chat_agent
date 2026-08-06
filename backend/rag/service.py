from __future__ import annotations

import csv
import hashlib
import html
import json
import math
import re
import shutil
import time
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup

from .bm25 import BM25Index
from .chunking import StructureFirstSemanticChunker, content_hash, tokenize
from .config import RagConfig, SemanticChunkingConfig, load_rag_config
from .models import DocumentManifest, KnowledgeChunk
from .retrieval import (
    BackendUnavailableError,
    EmbeddingProvider,
    QdrantVectorStore,
    RagBackendError,
    RerankerProvider,
    VectorRecord,
    VectorSearchHit,
    build_embedding_provider,
    build_reranker_provider,
    build_vector_store,
)
from .sources import (
    FeishuDocumentProvider,
    RemoteDocumentMetadata,
    RemoteDocumentProvider,
    RemoteDocumentSnapshot,
    RemoteSourceError,
)

_SUPPORTED_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".html", ".htm", ".pdf", ".docx"}
_PARSER_VERSION = "local-parser-v1"


def _canonical_json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)


def _doc_id(collection: str, source_uri: str) -> str:
    digest = hashlib.sha256(f"{collection}:{source_uri}".encode("utf-8")).hexdigest()[:16]
    return f"doc-{digest}"


def _safe_component(value: str, default: str = "default") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", str(value or "").strip()).strip("._-")
    return cleaned or default


def _safe_filename(value: str, default: str = "document") -> str:
    name = Path(value or default).name
    suffix = Path(name).suffix
    stem = name[: -len(suffix)] if suffix else name
    safe_stem = _safe_component(stem, default)
    safe_suffix = suffix.lower() if re.fullmatch(r"\.[A-Za-z0-9]+", suffix or "") else ""
    return f"{safe_stem}{safe_suffix}"


def _is_relative_to(path: Path, parent: Path) -> bool:
    try:
        path.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def _query_terms(query: str) -> list[str]:
    terms = {token for token in tokenize(query) if len(token) >= 2}
    for match in re.finditer(r"[\u4e00-\u9fff]{2,}", query or ""):
        value = match.group(0)
        max_size = min(12, len(value))
        for size in range(max_size, 1, -1):
            for start in range(0, len(value) - size + 1):
                terms.add(value[start:start + size])
    return sorted(terms, key=lambda value: (-len(value), value))


def _find_term_index(text: str, term: str) -> int:
    direct = text.lower().find(term.lower())
    if direct >= 0:
        return direct
    if not re.search(r"[\u4e00-\u9fff]", term):
        return -1
    compact_chars: list[str] = []
    original_indexes: list[int] = []
    for index, char in enumerate(text):
        if char.isspace():
            continue
        compact_chars.append(char)
        original_indexes.append(index)
    compact_index = "".join(compact_chars).lower().find(re.sub(r"\s+", "", term).lower())
    if compact_index < 0 or compact_index >= len(original_indexes):
        return -1
    return original_indexes[compact_index]


def _safe_snippet(text: str, query: str = "", limit: int = 360) -> str:
    normalized = re.sub(r"\s+", " ", text or "").strip()
    if len(normalized) <= limit:
        return normalized
    match_index = -1
    for term in _query_terms(query):
        match_index = _find_term_index(normalized, term)
        if match_index >= 0:
            break
    if match_index < 0:
        return normalized[: max(0, limit - 3)].rstrip() + "..."
    start = max(0, match_index - limit // 2)
    end = min(len(normalized), start + max(0, limit - 3))
    start = max(0, end - max(0, limit - 3))
    prefix = "..." if start > 0 else ""
    suffix = "..." if end < len(normalized) else ""
    return f"{prefix}{normalized[start:end].strip()}{suffix}"


def _safe_chunk_excerpt(text: str, query: str = "", limit: int = 1200) -> str:
    return _safe_snippet(text, query=query, limit=limit)


def _cosine(left: dict[str, float], right: dict[str, float]) -> float:
    if not left or not right:
        return 0.0
    dot = sum(value * right.get(key, 0.0) for key, value in left.items())
    left_norm = math.sqrt(sum(value * value for value in left.values()))
    right_norm = math.sqrt(sum(value * value for value in right.values()))
    if not left_norm or not right_norm:
        return 0.0
    return dot / (left_norm * right_norm)


def _vectorize(text: str) -> dict[str, float]:
    vector: dict[str, float] = {}
    for token in tokenize(text):
        vector[token] = vector.get(token, 0.0) + 1.0
    return vector


class PersonalKnowledgeBase:
    def __init__(
        self,
        *,
        workspace_root: str | Path,
        store_path: str | Path,
        config: RagConfig | None = None,
        chunker_config: SemanticChunkingConfig | None = None,
        reranker_enabled: bool | None = None,
        embedding_provider: EmbeddingProvider | None = None,
        reranker_provider: RerankerProvider | None = None,
        vector_store: Any | None = None,
        remote_provider: RemoteDocumentProvider | None = None,
    ) -> None:
        self.workspace_root = Path(workspace_root).resolve()
        self.store_path = Path(store_path).resolve()
        if config is None:
            config = load_rag_config(
                {
                    "enabled": True,
                    "retrieval_profile": "deterministic",
                    "knowledge_store_path": str(self.store_path),
                    "embedding_provider": "deterministic",
                    "vector_backend": "in_memory",
                    "sparse_backend": "in_memory_bm25",
                    "reranker": {
                        "enabled": bool(reranker_enabled),
                        "provider": "local_overlap",
                    },
                }
            )
        elif config.knowledge_store_path.resolve() != self.store_path:
            config = replace(config, knowledge_store_path=self.store_path)
        if chunker_config is not None:
            config = replace(config, chunking=chunker_config)
        if reranker_enabled is not None:
            config = replace(config, reranker=replace(config.reranker, enabled=bool(reranker_enabled)))
        self.config = config
        self.retrieval_profile = str(getattr(config, "retrieval_profile", "production") or "production")
        self.documents_path = self.store_path / "documents"
        self.index_path = self.store_path / "index"
        self.manifests_path = self.store_path / "manifests"
        self.sync_reports_path = self.store_path / "reports"
        for path in (self.store_path, self.documents_path, self.index_path, self.manifests_path, self.sync_reports_path):
            path.mkdir(parents=True, exist_ok=True)
        self.chunker_config = config.chunking
        self.chunker = StructureFirstSemanticChunker(self.chunker_config)
        self.reranker_enabled = bool(config.reranker.enabled)
        self.embedding_provider = embedding_provider
        self.reranker_provider = reranker_provider
        self.vector_store = vector_store
        self.remote_provider = remote_provider
        self._backend_error: RagBackendError | None = None
        self._initialize_retrieval_backends()
        self.manifests: dict[str, DocumentManifest] = {}
        self.chunks_by_doc: dict[str, list[KnowledgeChunk]] = {}
        self._load()

    @classmethod
    def from_config(cls, *, workspace_root: str | Path, config_overrides: dict[str, Any] | None = None) -> "PersonalKnowledgeBase":
        config = load_rag_config(config_overrides)
        return cls(
            workspace_root=workspace_root,
            store_path=config.knowledge_store_path,
            config=config,
        )

    def _initialize_retrieval_backends(self) -> None:
        """Prepare configured adapters without hiding initialization failures."""
        try:
            if self.embedding_provider is None:
                self.embedding_provider = build_embedding_provider(self.config)
            if self.config.vector_backend == "qdrant" and self.vector_store is None:
                self.vector_store = build_vector_store(self.config)
            if self.reranker_enabled and self.reranker_provider is None:
                self.reranker_provider = build_reranker_provider(self.config)
        except RagBackendError as exc:
            self._backend_error = exc

    def _error_payload(self, error: BaseException) -> dict[str, Any]:
        if isinstance(error, RagBackendError):
            return error.to_dict()
        return {"code": "rag_backend_error", "message": str(error)}

    def _stack_settings(
        self,
        *,
        latency_ms: float = 0.0,
        embedding_active: bool = False,
        vector_active: bool = False,
        reranker_active: bool = False,
        candidate_counts: dict[str, int] | None = None,
    ) -> dict[str, Any]:
        embedding_model = str(
            getattr(self.embedding_provider, "model_id", "") or self.config.embedding_model or ""
        )
        reranker_model = str(
            getattr(self.reranker_provider, "model_id", "") or self.config.reranker.model or ""
        )
        vector_accelerator = str(
            getattr(self.vector_store, "vector_accelerator", "none") or "none"
        )
        return {
            "retrieval_profile": self.retrieval_profile,
            "embedding_provider": self.config.embedding_provider,
            "embedding_model": embedding_model,
            "embedding_dimension": int(self.config.embedding_dimension),
            "embedding_active": bool(embedding_active),
            "vector_backend": self.config.vector_backend,
            "vector_collection": self.config.qdrant_collection,
            "vector_backend_active": bool(vector_active),
            "vector_accelerator": vector_accelerator,
            "vector_accelerator_enabled": bool(getattr(self.config, "faiss_enabled", False)),
            "vector_accelerator_active": bool(
                vector_accelerator != "none" and getattr(self.vector_store, "faiss_active", False)
            ),
            "sparse_backend": self.config.sparse_backend,
            "bm25_k1": round(float(self.config.bm25.k1), 6),
            "bm25_b": round(float(self.config.bm25.b), 6),
            "fusion_strategy": self.config.hybrid.fusion_strategy,
            "reranker_provider": self.config.reranker.provider,
            "reranker_model": reranker_model,
            "reranker_enabled": self.reranker_enabled,
            "reranker_active": bool(reranker_active),
            "candidate_counts": dict(candidate_counts or {}),
            "latency_ms": round(min(max(float(latency_ms), 0.0), 60000.0), 3),
        }

    def _backend_error_result(self, *, operation: str, error: BaseException, **payload: Any) -> dict[str, Any]:
        result = {
            **payload,
            "status": "error",
            "operation": operation,
            "error": self._error_payload(error),
            "settings": self._stack_settings(),
        }
        return result

    def _manifest_path(self) -> Path:
        return self.manifests_path / "manifests.json"

    def _chunks_path(self) -> Path:
        return self.index_path / "chunks.json"

    def _load(self) -> None:
        manifest_path = self._manifest_path()
        legacy_manifest_path = self.store_path / "manifests.json"
        if not manifest_path.exists() and legacy_manifest_path.exists():
            manifest_path = legacy_manifest_path
        if manifest_path.exists():
            data = json.loads(manifest_path.read_text(encoding="utf-8"))
            self.manifests = {
                doc_id: DocumentManifest(**item)
                for doc_id, item in data.items()
                if isinstance(item, dict)
            }
        chunks_path = self._chunks_path()
        legacy_chunks_path = self.store_path / "chunks.json"
        if not chunks_path.exists() and legacy_chunks_path.exists():
            chunks_path = legacy_chunks_path
        if chunks_path.exists():
            data = json.loads(chunks_path.read_text(encoding="utf-8"))
            self.chunks_by_doc = {
                doc_id: [KnowledgeChunk(**chunk) for chunk in chunks if isinstance(chunk, dict)]
                for doc_id, chunks in data.items()
                if isinstance(chunks, list)
            }

    def _save(self) -> None:
        for path in (self.store_path, self.documents_path, self.index_path, self.manifests_path, self.sync_reports_path):
            path.mkdir(parents=True, exist_ok=True)
        self._manifest_path().write_text(
            json.dumps({key: value.to_dict() for key, value in self.manifests.items()}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        self._chunks_path().write_text(
            json.dumps(
                {key: [chunk.to_dict() for chunk in chunks] for key, chunks in self.chunks_by_doc.items()},
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    def _validate_source(self, path: str | Path) -> tuple[Path | None, str]:
        candidate = Path(path).expanduser()
        resolved = candidate.resolve() if candidate.is_absolute() else (self.workspace_root / candidate).resolve()
        try:
            resolved.relative_to(self.workspace_root)
        except ValueError:
            return resolved, "outside allowed workspace roots"
        if not resolved.exists() or not resolved.is_file():
            return resolved, "file not found"
        if resolved.suffix.lower() not in _SUPPORTED_SUFFIXES:
            return resolved, "unsupported file type"
        return resolved, ""

    def _collection_from_document_path(self, path: Path) -> str:
        try:
            relative = path.resolve().relative_to(self.documents_path.resolve())
        except ValueError:
            return "default"
        if len(relative.parts) > 1:
            return _safe_component(relative.parts[0])
        return "default"

    def _target_import_path(self, collection: str, filename: str) -> Path:
        collection_dir = self.documents_path / _safe_component(collection)
        collection_dir.mkdir(parents=True, exist_ok=True)
        safe_name = _safe_filename(filename, "document")
        return collection_dir / safe_name

    def _read_supported_file(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            try:
                from pypdf import PdfReader
            except Exception as exc:
                raise RuntimeError("PDF text extraction requires pypdf") from exc
            with path.open("rb") as handle:
                reader = PdfReader(handle)
                pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(page.strip() for page in pages if page.strip()).strip()
            if not text:
                raise RuntimeError("PDF text extraction produced no text")
            return text
        if suffix == ".docx":
            try:
                from docx import Document
            except Exception as exc:
                raise RuntimeError("Word document text extraction requires python-docx") from exc
            doc = Document(str(path))
            paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)
            if not text:
                raise RuntimeError("Word document text extraction produced no text")
            return text
        if suffix == ".csv":
            with path.open("r", encoding="utf-8", newline="", errors="ignore") as handle:
                rows = list(csv.reader(handle))
            return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
        raw = path.read_text(encoding="utf-8", errors="ignore")
        if suffix in {".html", ".htm"}:
            return BeautifulSoup(raw, "html.parser").get_text("\n")
        return html.unescape(raw)

    def _retrieval_signature(self) -> str:
        return self.config.retrieval_signature()

    def _indexed_at(self) -> str:
        return datetime.now(timezone.utc).isoformat(timespec="seconds")

    def _document_vectors_current(self, doc_id: str) -> bool:
        if self.config.vector_backend != "qdrant":
            return True
        if self.vector_store is None:
            raise self._backend_error or BackendUnavailableError("Qdrant vector store is unavailable")
        expected_ids = {chunk.chunk_id for chunk in self.chunks_by_doc.get(doc_id, [])}
        if hasattr(self.vector_store, "document_point_ids"):
            observed_ids = set(self.vector_store.document_point_ids(doc_id))
            return observed_ids == expected_ids
        return bool(self.vector_store.document_exists(doc_id))

    def _vector_records(self, chunks: list[KnowledgeChunk], manifest: DocumentManifest) -> list[VectorRecord]:
        if self.config.vector_backend != "qdrant":
            return []
        if self.embedding_provider is None:
            raise self._backend_error or BackendUnavailableError("BGE-M3 embedding provider is unavailable")
        vectors = self.embedding_provider.embed([chunk.text for chunk in chunks])
        if len(vectors) != len(chunks):
            raise BackendUnavailableError(
                f"Embedding provider returned {len(vectors)} vectors for {len(chunks)} chunks"
            )
        records: list[VectorRecord] = []
        for chunk, vector in zip(chunks, vectors):
            metadata = dict(manifest.metadata)
            metadata.update(chunk.metadata)
            records.append(
                VectorRecord(
                    chunk_id=chunk.chunk_id,
                    doc_id=chunk.doc_id,
                    collection=chunk.collection,
                    vector=list(vector),
                    payload={
                        "source_ref": chunk.source_ref,
                        "source_type": manifest.source_type,
                        "title": manifest.title,
                        "ordinal": chunk.ordinal,
                        "heading_path": list(chunk.heading_path),
                        "metadata": metadata,
                    },
                )
            )
        return records

    def _commit_document(
        self,
        *,
        manifest: DocumentManifest,
        chunks: list[KnowledgeChunk],
    ) -> None:
        if self.config.vector_backend == "qdrant":
            if self.vector_store is None:
                raise self._backend_error or BackendUnavailableError("Qdrant vector store is unavailable")
            records = self._vector_records(chunks, manifest)
            self.vector_store.refresh(manifest.doc_id, records)
        self.manifests[manifest.doc_id] = manifest
        self.chunks_by_doc[manifest.doc_id] = chunks

    @staticmethod
    def _is_remote_manifest(manifest: DocumentManifest) -> bool:
        return bool(manifest.metadata.get("remote_provider"))

    def _remote_snapshot_path(self, collection: str, doc_id: str) -> Path:
        target = self.documents_path / "_remote" / _safe_component(collection) / f"{_safe_component(doc_id)}.md"
        target.parent.mkdir(parents=True, exist_ok=True)
        return target

    def _manifest_snapshot_path(self, manifest: DocumentManifest) -> Path | None:
        snapshot_value = str(manifest.metadata.get("snapshot_path") or "").strip()
        if not snapshot_value:
            return None
        candidate = Path(snapshot_value)
        if not candidate.is_absolute():
            candidate = self.store_path / candidate
        candidate = candidate.resolve()
        remote_root = (self.documents_path / "_remote").resolve()
        return candidate if _is_relative_to(candidate, remote_root) else None

    def _is_remote_snapshot_path(self, path: Path) -> bool:
        return _is_relative_to(path, self.documents_path / "_remote")

    def _remote_provider_for(self, provider_name: str | None = None) -> RemoteDocumentProvider:
        selected = str(provider_name or self.config.remote_source.provider or "feishu").strip().lower()
        if not self.config.remote_source.enabled:
            raise RemoteSourceError(
                "Remote document sources are disabled",
                category="provider_unavailable",
                retryable=False,
            )
        if self.remote_provider is not None:
            actual = str(getattr(self.remote_provider, "provider_name", selected) or selected).lower()
            if actual != selected:
                raise RemoteSourceError(
                    f"Configured remote provider does not match {selected}",
                    category="provider_unavailable",
                    retryable=False,
                )
            return self.remote_provider
        if selected == "feishu":
            return FeishuDocumentProvider(
                workspace_root=self.workspace_root,
                session_file=self.config.remote_source.session_file,
                max_content_chars=self.config.remote_source.max_content_chars,
            )
        raise RemoteSourceError(
            f"Unsupported remote document provider: {selected}",
            category="provider_unavailable",
            retryable=False,
        )

    def _remote_call(self, operation: str, callback: Any) -> Any:
        max_retries = max(0, int(self.config.remote_source.max_retries))
        backoff = max(0.0, float(self.config.remote_source.retry_backoff_seconds))
        for attempt in range(max_retries + 1):
            try:
                return callback()
            except RemoteSourceError as exc:
                if not exc.retryable or attempt >= max_retries:
                    raise
                delay = min(backoff * (2**attempt), 5.0)
                if delay > 0:
                    time.sleep(delay)
            except Exception as exc:
                raise RemoteSourceError(
                    f"Remote provider failed during {operation}",
                    category="provider_unavailable",
                    retryable=True,
                ) from exc
        raise RemoteSourceError(
            f"Remote provider failed during {operation}",
            category="provider_unavailable",
            retryable=True,
        )

    def _remote_result(self, status: str, manifest: DocumentManifest) -> dict[str, Any]:
        metadata = manifest.metadata
        return {
            "status": status,
            "collection": manifest.collection,
            "doc_id": manifest.doc_id,
            "source_uri": manifest.source_uri,
            "source_url": str(metadata.get("source_url") or ""),
            "source_type": manifest.source_type,
            "title": manifest.title,
            "chunk_count": manifest.chunk_count,
            "content_hash": manifest.content_hash,
            "remote_revision": str(metadata.get("remote_revision") or ""),
            "remote_updated_at": str(metadata.get("remote_updated_at") or ""),
            "owner": str(metadata.get("owner") or ""),
            "content_format": str(metadata.get("content_format") or "markdown"),
            "snapshot_path": str(metadata.get("snapshot_path") or ""),
            "settings": self._stack_settings(
                embedding_active=bool(self.config.vector_backend == "qdrant" and manifest.chunk_count),
                vector_active=bool(self.vector_store is not None and manifest.chunk_count),
            ),
        }

    def _write_remote_snapshot(self, path: Path, content: str) -> None:
        temporary = path.with_name(f".{path.name}.tmp")
        try:
            temporary.write_text(content, encoding="utf-8")
            temporary.replace(path)
        finally:
            temporary.unlink(missing_ok=True)

    def _remote_metadata_dict(
        self,
        metadata: RemoteDocumentMetadata,
        *,
        snapshot_path: Path,
        content_format: str,
        existing: DocumentManifest | None = None,
    ) -> dict[str, Any]:
        result = dict(existing.metadata if existing else {})
        result.update(
            {
                "remote_provider": metadata.provider,
                "object_token": metadata.object_token,
                "source_url": metadata.source_url,
                "owner": metadata.owner,
                "remote_revision": metadata.remote_revision,
                "remote_updated_at": metadata.remote_updated_at,
                "content_format": str(content_format or "markdown").strip().lower()[:32] or "markdown",
                "snapshot_path": str(snapshot_path),
                "active": True,
            }
        )
        if metadata.extra:
            result["provider_metadata"] = dict(metadata.extra)
        return result

    def _index_remote_snapshot(
        self,
        collection: str,
        snapshot: RemoteDocumentSnapshot,
        *,
        refresh: bool = False,
    ) -> dict[str, Any]:
        normalized_collection = _safe_component(collection)
        metadata = snapshot.metadata
        source_uri = str(metadata.source_uri or "").strip()
        if not source_uri:
            raise RemoteSourceError("Remote source URI is required", category="invalid_reference")
        content = str(snapshot.content or "")
        max_content_chars = max(1, int(self.config.remote_source.max_content_chars))
        if not content.strip():
            raise RemoteSourceError(
                "Remote document contains no indexable text",
                category="content_invalid",
                retryable=False,
            )
        if len(content) > max_content_chars or len(content.encode("utf-8")) > max_content_chars:
            raise RemoteSourceError(
                "Remote document exceeds the configured content limit",
                category="content_invalid",
                retryable=False,
                details={"max_content_chars": max_content_chars},
            )
        doc_id = _doc_id(normalized_collection, source_uri)
        existing = self.manifests.get(doc_id)
        retrieval_signature = self._retrieval_signature()
        chunk_signature = self.chunker_config.signature()
        same_revision = bool(
            metadata.remote_revision
            and existing is not None
            and existing.metadata.get("remote_revision") == metadata.remote_revision
        )
        unchanged = bool(
            existing is not None
            and not refresh
            and existing.status == "indexed"
            and existing.metadata.get("active", True)
            and existing.content_hash == snapshot.content_hash
            and existing.chunking_signature == chunk_signature
            and existing.retrieval_signature == retrieval_signature
            and self._document_vectors_current(doc_id)
            and (not metadata.remote_revision or same_revision)
        )
        if unchanged:
            return self._remote_result("unchanged", existing)

        chunks = self.chunker.chunk_text(
            snapshot.content,
            source_ref=metadata.title or source_uri,
            doc_id=doc_id,
            collection=normalized_collection,
        )
        snapshot_path = self._remote_snapshot_path(normalized_collection, doc_id)
        manifest = DocumentManifest(
            doc_id=doc_id,
            collection=normalized_collection,
            source_uri=source_uri,
            source_type=f"{metadata.provider}_document",
            content_hash=snapshot.content_hash,
            parser_version=_PARSER_VERSION,
            chunker_version=self.chunker_config.chunker_version,
            chunking_signature=chunk_signature,
            title=metadata.title or source_uri,
            status="indexed",
            chunk_count=len(chunks),
            metadata=self._remote_metadata_dict(
                metadata,
                snapshot_path=snapshot_path,
                content_format=snapshot.content_format,
                existing=existing,
            ),
            retrieval_signature=retrieval_signature,
            indexed_at=self._indexed_at(),
        )
        previous_manifest = existing
        previous_chunks = list(self.chunks_by_doc.get(doc_id, []))
        previous_snapshot_path = self._manifest_snapshot_path(existing) if existing else None
        previous_snapshot = None
        if previous_snapshot_path is not None and previous_snapshot_path.exists():
            previous_snapshot = previous_snapshot_path.read_bytes()
        try:
            self._write_remote_snapshot(snapshot_path, snapshot.content)
            self._commit_document(manifest=manifest, chunks=chunks)
            self._save()
        except Exception:
            try:
                if previous_snapshot is not None and previous_snapshot_path is not None:
                    previous_snapshot_path.parent.mkdir(parents=True, exist_ok=True)
                    previous_snapshot_path.write_bytes(previous_snapshot)
                elif snapshot_path.exists():
                    snapshot_path.unlink(missing_ok=True)
            except OSError:
                pass
            try:
                if self.config.vector_backend == "qdrant" and self.vector_store is not None:
                    if previous_manifest is not None:
                        self.vector_store.refresh(
                            doc_id,
                            self._vector_records(previous_chunks, previous_manifest),
                        )
                    else:
                        self.vector_store.delete_document(doc_id)
            except Exception:
                pass
            if previous_manifest is None:
                self.manifests.pop(doc_id, None)
                self.chunks_by_doc.pop(doc_id, None)
            else:
                self.manifests[doc_id] = previous_manifest
                self.chunks_by_doc[doc_id] = previous_chunks
            raise
        status = "refreshed" if existing is not None else "indexed"
        return self._remote_result(status, manifest)

    def import_remote_document(
        self,
        collection: str,
        reference: str,
        *,
        refresh: bool = False,
        provider_name: str | None = None,
    ) -> dict[str, Any]:
        started = time.perf_counter()
        normalized_collection = _safe_component(collection)
        existing_doc_id = ""
        try:
            provider = self._remote_provider_for(provider_name)
            metadata = self._remote_call("inspect", lambda: provider.inspect(reference))
            existing_doc_id = _doc_id(normalized_collection, metadata.source_uri)
            existing = self.manifests.get(existing_doc_id)
            if (
                existing is not None
                and not refresh
                and existing.status == "indexed"
                and existing.metadata.get("active", True)
                and metadata.remote_revision
                and existing.metadata.get("remote_revision") == metadata.remote_revision
                and existing.chunking_signature == self.chunker_config.signature()
                and existing.retrieval_signature == self._retrieval_signature()
                and self._document_vectors_current(existing_doc_id)
            ):
                result = self._remote_result("unchanged", existing)
                result["settings"]["latency_ms"] = round((time.perf_counter() - started) * 1000.0, 3)
                return result
            snapshot = self._remote_call("fetch", lambda: provider.fetch(metadata))
            result = self._index_remote_snapshot(normalized_collection, snapshot, refresh=refresh)
            result["settings"]["latency_ms"] = round((time.perf_counter() - started) * 1000.0, 3)
            return result
        except (RemoteSourceError, RagBackendError) as exc:
            return self._backend_error_result(
                operation="import_remote_document",
                error=exc,
                collection=normalized_collection,
                doc_id=existing_doc_id,
            )
        except Exception:
            return self._backend_error_result(
                operation="import_remote_document",
                error=RagBackendError("Remote document indexing failed."),
                collection=normalized_collection,
                doc_id=existing_doc_id,
            )

    def _mark_remote_sync_failed(self, manifest: DocumentManifest, error: BaseException) -> None:
        manifest.status = "sync_failed"
        manifest.latest_error = json.dumps(self._error_payload(error), ensure_ascii=False)
        manifest.metadata["active"] = True

    def _deactivate_remote_manifest(self, manifest: DocumentManifest, error: BaseException, status: str) -> None:
        self._delete_document_vectors(manifest.doc_id)
        self.manifests[manifest.doc_id] = manifest
        self.chunks_by_doc.pop(manifest.doc_id, None)
        snapshot_path = self._manifest_snapshot_path(manifest)
        if snapshot_path is not None:
            snapshot_path.unlink(missing_ok=True)
        manifest.status = status
        manifest.chunk_count = 0
        manifest.latest_error = json.dumps(self._error_payload(error), ensure_ascii=False)
        manifest.metadata["active"] = False
        manifest.metadata.pop("snapshot_path", None)

    def sync_remote_sources(
        self,
        *,
        provider_name: str | None = None,
        collection: str | None = None,
        doc_ids: list[str] | None = None,
        dry_run: bool = False,
    ) -> dict[str, Any]:
        started = time.perf_counter()
        selected_provider = str(provider_name or self.config.remote_source.provider or "feishu").lower()
        normalized_collection = _safe_component(collection) if collection else None
        selected_ids = set(doc_ids or [])
        counts = {
            "indexed": 0,
            "refreshed": 0,
            "unchanged": 0,
            "failed": 0,
            "access_denied": 0,
            "remote_deleted": 0,
            "deleted": 0,
            "would_refresh": 0,
            "would_access_denied": 0,
            "would_delete": 0,
        }
        files: list[dict[str, Any]] = []
        try:
            provider = self._remote_provider_for(selected_provider)
        except RemoteSourceError as exc:
            error = self._error_payload(exc)
            return {
                "provider": selected_provider,
                "collection": normalized_collection,
                "dry_run": dry_run,
                "counts": {**counts, "failed": 1},
                "files": [{"status": "sync_failed", "error": error}],
                "error": error,
                "settings": self._stack_settings(
                    latency_ms=(time.perf_counter() - started) * 1000.0,
                    candidate_counts={**counts, "failed": 1},
                ),
            }
        manifests = [
            manifest
            for manifest in self.manifests.values()
            if manifest.metadata.get("remote_provider") == selected_provider
            and (not normalized_collection or manifest.collection == normalized_collection)
            and (not selected_ids or manifest.doc_id in selected_ids)
        ]
        for manifest in sorted(manifests, key=lambda item: item.source_uri):
            reference = str(manifest.metadata.get("source_url") or manifest.metadata.get("object_token") or manifest.source_uri)
            try:
                metadata = self._remote_call("inspect", lambda: provider.inspect(reference))
                same_revision = bool(
                    metadata.remote_revision
                    and manifest.metadata.get("remote_revision") == metadata.remote_revision
                )
                vectors_current = self._document_vectors_current(manifest.doc_id)
                if same_revision and manifest.status == "indexed" and manifest.metadata.get("active", True) and vectors_current:
                    counts["unchanged"] += 1
                    files.append({"doc_id": manifest.doc_id, "source_uri": manifest.source_uri, "status": "unchanged"})
                    continue
                if dry_run:
                    files.append({
                        "doc_id": manifest.doc_id,
                        "source_uri": manifest.source_uri,
                        "status": "would_refresh",
                    })
                    counts["would_refresh"] += 1
                    continue
                snapshot = self._remote_call("fetch", lambda: provider.fetch(metadata))
                result = self._index_remote_snapshot(manifest.collection, snapshot, refresh=True)
                status = result["status"]
                counts[status] = counts.get(status, 0) + 1
                files.append({"doc_id": manifest.doc_id, "source_uri": manifest.source_uri, "status": status})
            except RemoteSourceError as exc:
                if exc.category in {"permission_denied", "not_found"}:
                    is_deleted = exc.category == "not_found"
                    if not dry_run:
                        self._deactivate_remote_manifest(
                            manifest,
                            exc,
                            "access_denied" if not is_deleted else "remote_deleted",
                        )
                    if dry_run:
                        status = "would_access_denied" if not is_deleted else "would_delete"
                        counts[status] += 1
                    else:
                        status = "access_denied" if not is_deleted else "remote_deleted"
                        counts[status] += 1
                        if is_deleted:
                            counts["deleted"] += 1
                    files.append({
                        "doc_id": manifest.doc_id,
                        "source_uri": manifest.source_uri,
                        "status": status,
                        "error": self._error_payload(exc),
                    })
                else:
                    if not dry_run:
                        self._mark_remote_sync_failed(manifest, exc)
                    counts["failed"] += 1
                    files.append({
                        "doc_id": manifest.doc_id,
                        "source_uri": manifest.source_uri,
                        "status": "sync_failed",
                        "error": self._error_payload(exc),
                    })
            except RagBackendError as exc:
                if not dry_run:
                    self._mark_remote_sync_failed(manifest, exc)
                counts["failed"] += 1
                files.append({
                    "doc_id": manifest.doc_id,
                    "source_uri": manifest.source_uri,
                    "status": "sync_failed",
                    "error": self._error_payload(exc),
                })
        if not dry_run:
            self._save()
        report = {
            "provider": selected_provider,
            "collection": normalized_collection,
            "dry_run": dry_run,
            "counts": counts,
            "files": files,
            "settings": self._stack_settings(
                latency_ms=(time.perf_counter() - started) * 1000.0,
                candidate_counts=dict(counts),
            ),
        }
        errors = [item.get("error") for item in files if isinstance(item.get("error"), dict)]
        if errors:
            report["error"] = errors[0]
        return report

    def index_files(
        self,
        collection: str,
        paths: list[str | Path],
        *,
        refresh: bool = False,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        started = time.perf_counter()
        indexed: list[dict[str, Any]] = []
        skipped: list[dict[str, Any]] = []
        failed: list[dict[str, Any]] = []
        normalized_collection = _safe_component(collection)
        retrieval_signature = self._retrieval_signature()
        for path in paths:
            resolved, reason = self._validate_source(path)
            if reason:
                skipped.append({"path": str(resolved or path), "reason": reason})
                continue
            assert resolved is not None
            source_uri = str(resolved)
            text = self._read_supported_file(resolved)
            digest = content_hash(text)
            doc_id = _doc_id(normalized_collection, source_uri)
            existing = self.manifests.get(doc_id)
            chunk_signature = self.chunker_config.signature()
            needs_refresh = (
                refresh
                or existing is None
                or existing.content_hash != digest
                or existing.chunking_signature != chunk_signature
                or existing.retrieval_signature != retrieval_signature
            )
            try:
                if existing is not None and not needs_refresh:
                    if not self._document_vectors_current(doc_id):
                        needs_refresh = True
                if not needs_refresh:
                    indexed.append(
                        {
                            "doc_id": doc_id,
                            "source_uri": source_uri,
                            "refreshed": False,
                            "chunk_count": existing.chunk_count,
                            "status": "unchanged",
                        }
                    )
                    continue
                chunks = self.chunker.chunk_text(
                    text,
                    source_ref=resolved.name,
                    doc_id=doc_id,
                    collection=normalized_collection,
                )
                manifest = DocumentManifest(
                    doc_id=doc_id,
                    collection=normalized_collection,
                    source_uri=source_uri,
                    source_type=resolved.suffix.lower().lstrip("."),
                    content_hash=digest,
                    parser_version=_PARSER_VERSION,
                    chunker_version=self.chunker_config.chunker_version,
                    chunking_signature=chunk_signature,
                    title=resolved.name,
                    status="indexed",
                    chunk_count=len(chunks),
                    metadata=dict(metadata or (existing.metadata if existing else {})),
                    retrieval_signature=retrieval_signature,
                    indexed_at=self._indexed_at(),
                )
                self._commit_document(manifest=manifest, chunks=chunks)
                indexed.append(
                    {
                        "doc_id": doc_id,
                        "source_uri": source_uri,
                        "refreshed": existing is not None,
                        "chunk_count": len(chunks),
                        "status": "indexed",
                    }
                )
            except RagBackendError as exc:
                failure = {
                    "path": source_uri,
                    "doc_id": doc_id,
                    "status": "failed",
                    "error": exc.to_dict(),
                }
                failed.append(failure)
            except Exception as exc:
                failure = {
                    "path": source_uri,
                    "doc_id": doc_id,
                    "status": "failed",
                    "error": self._error_payload(exc),
                }
                failed.append(failure)
        self._save()
        result = {
            "collection": normalized_collection,
            "indexed": indexed,
            "skipped": skipped,
            "failed": failed,
            "settings": self._stack_settings(
                latency_ms=(time.perf_counter() - started) * 1000.0,
                embedding_active=bool(indexed and self.config.vector_backend == "qdrant"),
                vector_active=bool(indexed and self.vector_store is not None),
                candidate_counts={"indexed": len(indexed), "skipped": len(skipped), "failed": len(failed)},
            ),
        }
        if failed:
            result["error"] = failed[0]["error"]
        return result

    def import_files(
        self,
        collection: str,
        paths: list[str | Path],
        *,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        del metadata
        normalized_collection = _safe_component(collection)
        imported: list[dict[str, Any]] = []
        skipped: list[dict[str, Any]] = []
        for path in paths:
            resolved, reason = self._validate_source(path)
            if reason:
                skipped.append({"path": str(resolved or path), "reason": reason})
                continue
            assert resolved is not None
            if _is_relative_to(resolved, self.documents_path):
                target = resolved
            else:
                target = self._target_import_path(normalized_collection, resolved.name)
                if resolved.resolve() != target.resolve():
                    shutil.copy2(resolved, target)
            imported.append(
                {
                    "collection": normalized_collection,
                    "source_uri": str(resolved),
                    "knowledge_path": str(target),
                    "title": target.name,
                    "status": "pending_sync",
                }
            )
        return {
            "collection": normalized_collection,
            "imported": imported,
            "skipped": skipped,
            "counts": {
                "imported": len(imported),
                "skipped": len(skipped),
            },
        }

    def import_file_bytes(
        self,
        collection: str,
        filename: str,
        content: bytes,
        *,
        metadata: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        del metadata
        normalized_collection = _safe_component(collection)
        target = self._target_import_path(normalized_collection, filename)
        if target.suffix.lower() not in _SUPPORTED_SUFFIXES:
            return {
                "collection": normalized_collection,
                "imported": [],
                "skipped": [{"path": str(target), "reason": "unsupported file type"}],
                "counts": {"imported": 0, "skipped": 1},
            }
        target.write_bytes(content)
        return {
            "collection": normalized_collection,
            "imported": [
                {
                    "collection": normalized_collection,
                    "source_uri": str(target),
                    "knowledge_path": str(target),
                    "title": target.name,
                    "status": "pending_sync",
                }
            ],
            "skipped": [],
            "counts": {"imported": 1, "skipped": 0},
        }

    def sync(self, collection: str | None = None, *, dry_run: bool = False) -> dict[str, Any]:
        started = time.perf_counter()
        normalized_collection = _safe_component(collection) if collection else None
        counts = {
            "indexed": 0,
            "refreshed": 0,
            "unchanged": 0,
            "skipped": 0,
            "failed": 0,
            "deleted": 0,
        }
        files: list[dict[str, Any]] = []
        supported_paths: set[str] = set()
        scan_root = self.documents_path / normalized_collection if normalized_collection else self.documents_path
        candidates = sorted(scan_root.rglob("*")) if scan_root.exists() else []
        chunk_signature = self.chunker_config.signature()
        retrieval_signature = self._retrieval_signature()

        for path in candidates:
            if not path.is_file():
                continue
            if self._is_remote_snapshot_path(path):
                continue
            inferred_collection = self._collection_from_document_path(path)
            if normalized_collection and inferred_collection != normalized_collection:
                continue
            if path.suffix.lower() not in _SUPPORTED_SUFFIXES:
                counts["skipped"] += 1
                files.append({"path": str(path), "collection": inferred_collection, "status": "skipped", "reason": "unsupported file type"})
                continue
            source_uri = str(path.resolve())
            supported_paths.add(source_uri)
            doc_id = _doc_id(inferred_collection, source_uri)
            existing = self.manifests.get(doc_id)
            try:
                text = self._read_supported_file(path)
                digest = content_hash(text)
                needs_index = (
                    existing is None
                    or existing.content_hash != digest
                    or existing.chunking_signature != chunk_signature
                    or existing.retrieval_signature != retrieval_signature
                )
                if existing is not None and not needs_index:
                    needs_index = not self._document_vectors_current(doc_id)
                if not needs_index:
                    counts["unchanged"] += 1
                    files.append({"path": source_uri, "collection": inferred_collection, "doc_id": doc_id, "status": "unchanged"})
                    continue
                if dry_run:
                    status = "would_index" if existing is None else "would_refresh"
                    files.append({"path": source_uri, "collection": inferred_collection, "doc_id": doc_id, "status": status})
                    continue
                result = self.index_files(inferred_collection, [path], refresh=True)
                if result["indexed"]:
                    item = result["indexed"][0]
                    if item.get("refreshed"):
                        counts["refreshed"] += 1
                        status = "refreshed"
                    else:
                        counts["indexed"] += 1
                        status = "indexed"
                    files.append({"path": source_uri, "collection": inferred_collection, "doc_id": item["doc_id"], "status": status})
                elif result.get("failed"):
                    failure = result["failed"][0]
                    counts["failed"] += 1
                    files.append(
                        {
                            "path": source_uri,
                            "collection": inferred_collection,
                            "doc_id": doc_id,
                            "status": "failed",
                            "error": failure.get("error", {}),
                        }
                    )
                for skipped in result.get("skipped", []):
                    counts["skipped"] += 1
                    files.append({**skipped, "collection": inferred_collection, "status": "skipped"})
            except Exception as exc:
                counts["failed"] += 1
                files.append(
                    {
                        "path": source_uri,
                        "collection": inferred_collection,
                        "doc_id": doc_id,
                        "status": "failed",
                        "error": self._error_payload(exc),
                    }
                )

        deleted: list[str] = []
        for doc_id, manifest in list(self.manifests.items()):
            if self._is_remote_manifest(manifest):
                continue
            source_path = Path(manifest.source_uri)
            if normalized_collection and manifest.collection != normalized_collection:
                continue
            if normalized_collection and not _is_relative_to(source_path, scan_root):
                continue
            if not _is_relative_to(source_path, self.documents_path):
                continue
            if str(source_path.resolve()) in supported_paths:
                continue
            deleted.append(doc_id)

        if not dry_run:
            for doc_id in deleted:
                delete_result = self.delete_document(doc_id=doc_id)
                if delete_result.get("error"):
                    counts["failed"] += 1
                    files.append(
                        {
                            "doc_id": doc_id,
                            "status": "failed",
                            "error": delete_result.get("error"),
                        }
                    )
        counts["deleted"] = len(deleted)
        for doc_id in deleted:
            files.append({"doc_id": doc_id, "status": "deleted"})

        report = {
            "collection": normalized_collection,
            "dry_run": dry_run,
            "counts": counts,
            "files": files,
            "store_path": str(self.store_path),
            "settings": self._stack_settings(
                latency_ms=(time.perf_counter() - started) * 1000.0,
                embedding_active=bool(counts["indexed"] or counts["refreshed"]),
                vector_active=bool(self.vector_store is not None and (counts["indexed"] or counts["refreshed"])),
                candidate_counts=dict(counts),
            ),
        }
        if not dry_run:
            report_path = self.sync_reports_path / f"sync-{int(time.time() * 1000)}.json"
            report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            report["report_path"] = str(report_path)
        failures = [item.get("error") for item in files if isinstance(item.get("error"), dict)]
        if failures:
            report["error"] = failures[0]
        return report

    def list_collections(self) -> list[dict[str, Any]]:
        totals: dict[str, dict[str, Any]] = {}
        for manifest in self.manifests.values():
            entry = totals.setdefault(manifest.collection, {"collection": manifest.collection, "document_count": 0, "chunk_count": 0})
            entry["document_count"] += 1
            entry["chunk_count"] += manifest.chunk_count
        return sorted(totals.values(), key=lambda item: item["collection"])

    def list_documents(self, collection: str | None = None) -> list[dict[str, Any]]:
        return [
            manifest.to_dict()
            for manifest in sorted(self.manifests.values(), key=lambda item: item.source_uri)
            if collection is None or manifest.collection == collection
        ]

    def list_source_files(self, collection: str | None = None) -> list[dict[str, Any]]:
        normalized_collection = _safe_component(collection) if collection else None
        scan_root = self.documents_path / normalized_collection if normalized_collection else self.documents_path
        if not scan_root.exists():
            return []
        by_source = {manifest.source_uri: manifest for manifest in self.manifests.values()}
        sources: list[dict[str, Any]] = []
        for path in sorted(scan_root.rglob("*")):
            if not path.is_file():
                continue
            if self._is_remote_snapshot_path(path):
                continue
            inferred_collection = self._collection_from_document_path(path)
            if normalized_collection and inferred_collection != normalized_collection:
                continue
            source_uri = str(path.resolve())
            suffix = path.suffix.lower()
            manifest = by_source.get(source_uri)
            status = "pending_sync"
            reason = ""
            chunk_count = 0
            doc_id = ""
            source_type = suffix.lstrip(".")
            latest_error = ""
            if suffix not in _SUPPORTED_SUFFIXES:
                status = "unsupported"
                reason = "unsupported file type"
            elif manifest is not None:
                doc_id = manifest.doc_id
                chunk_count = manifest.chunk_count
                source_type = manifest.source_type
                latest_error = manifest.latest_error
                try:
                    current_content = content_hash(self._read_supported_file(path)) == manifest.content_hash
                    current_signature = manifest.retrieval_signature == self._retrieval_signature()
                    current_vectors = self._document_vectors_current(manifest.doc_id)
                    status = "indexed" if current_content and current_signature and current_vectors else "pending_sync"
                except Exception as exc:
                    status = "failed"
                    reason = str(exc)
            sources.append(
                {
                    "collection": inferred_collection,
                    "doc_id": doc_id,
                    "source_uri": source_uri,
                    "title": path.name,
                    "source_type": source_type,
                    "status": status,
                    "chunk_count": chunk_count,
                    "reason": reason or latest_error,
                }
            )
        for manifest in sorted(self.manifests.values(), key=lambda item: item.source_uri):
            if not self._is_remote_manifest(manifest):
                continue
            if normalized_collection and manifest.collection != normalized_collection:
                continue
            sources.append(
                {
                    "collection": manifest.collection,
                    "doc_id": manifest.doc_id,
                    "source_uri": manifest.source_uri,
                    "source_url": manifest.metadata.get("source_url", ""),
                    "title": manifest.title,
                    "source_type": manifest.source_type,
                    "status": manifest.status,
                    "chunk_count": manifest.chunk_count,
                    "reason": manifest.latest_error,
                }
            )
        return sources

    def get_document_detail(self, doc_id: str) -> dict[str, Any] | None:
        manifest = self.manifests.get(doc_id)
        if manifest is None:
            return None
        return {
            "document": manifest.to_dict(),
            "chunks": [chunk.to_dict() for chunk in self.chunks_by_doc.get(doc_id, [])],
        }

    def list_chunks(self, doc_id: str) -> list[dict[str, Any]]:
        return [chunk.to_dict() for chunk in self.chunks_by_doc.get(doc_id, [])]

    def _delete_document_vectors(self, doc_id: str) -> None:
        if self.config.vector_backend != "qdrant":
            return
        if self.vector_store is None:
            raise self._backend_error or BackendUnavailableError("Qdrant vector store is unavailable")
        self.vector_store.delete_document(doc_id)

    def delete_source_file(self, source_uri: str) -> dict[str, Any]:
        remote_matches = [
            manifest.doc_id
            for manifest in self.manifests.values()
            if self._is_remote_manifest(manifest) and manifest.source_uri == str(source_uri).strip()
        ]
        if remote_matches:
            return self.delete_document(doc_id=remote_matches[0], remove_source=True)
        if "://" in str(source_uri):
            return {"deleted": [], "removed_sources": [], "error": "remote source not found"}
        source_path = Path(source_uri)
        if not source_path.is_absolute():
            source_path = self.workspace_root / source_path
        resolved = source_path.resolve()
        if not _is_relative_to(resolved, self.documents_path):
            return {"deleted": [], "removed_sources": [], "error": "source is outside knowledge documents"}
        if not resolved.exists() or not resolved.is_file():
            return {"deleted": [], "removed_sources": [], "error": "source file not found"}

        deleted: list[str] = []
        for manifest_id, manifest in list(self.manifests.items()):
            if Path(manifest.source_uri).resolve() != resolved:
                continue
            try:
                self._delete_document_vectors(manifest_id)
            except RagBackendError as exc:
                return self._backend_error_result(
                    operation="delete_source_file",
                    error=exc,
                    deleted=[],
                    removed_sources=[],
                )
            self.manifests.pop(manifest_id, None)
            self.chunks_by_doc.pop(manifest_id, None)
            deleted.append(manifest_id)

        resolved.unlink()
        self._save()
        return {"deleted": deleted, "removed_sources": [str(resolved)]}

    def delete_document(
        self,
        *,
        doc_id: str | None = None,
        collection: str | None = None,
        source_uri: str | None = None,
        remove_source: bool = False,
    ) -> dict[str, Any]:
        to_delete: list[str] = []
        for manifest_id, manifest in self.manifests.items():
            if doc_id and manifest_id != doc_id:
                continue
            if collection and manifest.collection != collection:
                continue
            if source_uri:
                requested_source = str(source_uri).strip()
                if "://" in requested_source:
                    if manifest.source_uri != requested_source:
                        continue
                elif manifest.source_uri != str(Path(requested_source).resolve()):
                    continue
            to_delete.append(manifest_id)
        removed_sources: list[str] = []
        for manifest_id in to_delete:
            manifest = self.manifests.get(manifest_id)
            try:
                self._delete_document_vectors(manifest_id)
            except RagBackendError as exc:
                return self._backend_error_result(
                    operation="delete_document",
                    error=exc,
                    deleted=[],
                    removed_sources=[],
                )
            self.manifests.pop(manifest_id, None)
            self.chunks_by_doc.pop(manifest_id, None)
            if remove_source and manifest:
                if self._is_remote_manifest(manifest):
                    snapshot_path = self._manifest_snapshot_path(manifest)
                    if snapshot_path is not None:
                        snapshot_path.unlink(missing_ok=True)
                        removed_sources.append(str(snapshot_path))
                else:
                    source_path = Path(manifest.source_uri)
                    if _is_relative_to(source_path, self.documents_path):
                        source_path.unlink(missing_ok=True)
                        removed_sources.append(str(source_path))
        self._save()
        return {"deleted": to_delete, "removed_sources": removed_sources}

    @staticmethod
    def _filter_value_matches(actual: Any, expected: Any) -> bool:
        if isinstance(expected, (list, tuple, set)):
            expected_values = set(expected)
            if isinstance(actual, (list, tuple, set)):
                return bool(expected_values.intersection(actual))
            return actual in expected_values
        if isinstance(actual, (list, tuple, set)):
            return expected in actual
        return actual == expected

    def _chunk_matches_filters(
        self,
        chunk: KnowledgeChunk,
        manifest: DocumentManifest,
        filters: dict[str, Any] | None,
    ) -> bool:
        for raw_key, expected in (filters or {}).items():
            if expected is None:
                continue
            key = str(raw_key)
            if key.startswith("metadata."):
                key = key[len("metadata.") :]
            if key in {"doc_id", "collection", "source_uri", "source_type", "title"}:
                actual = {
                    "doc_id": manifest.doc_id,
                    "collection": manifest.collection,
                    "source_uri": manifest.source_uri,
                    "source_type": manifest.source_type,
                    "title": manifest.title,
                }[key]
            elif key in chunk.metadata:
                actual = chunk.metadata.get(key)
            else:
                actual = manifest.metadata.get(key)
            if not self._filter_value_matches(actual, expected):
                return False
        return True

    def _candidate_chunks(
        self,
        collection: str | None,
        filters: dict[str, Any] | None = None,
    ) -> list[KnowledgeChunk]:
        chunks: list[KnowledgeChunk] = []
        for doc_id, doc_chunks in self.chunks_by_doc.items():
            manifest = self.manifests.get(doc_id)
            if manifest is None:
                continue
            if self._is_remote_manifest(manifest) and not manifest.metadata.get("active", True):
                continue
            if collection and manifest.collection != collection:
                continue
            chunks.extend(
                chunk
                for chunk in doc_chunks
                if self._chunk_matches_filters(chunk, manifest, filters)
            )
        return chunks

    def _bm25_rank(self, query_tokens: list[str], chunks: list[KnowledgeChunk]) -> list[tuple[KnowledgeChunk, float]]:
        if not query_tokens or not chunks:
            return []
        index = BM25Index(
            chunks,
            k1=self.config.bm25.k1,
            b=self.config.bm25.b,
        )
        depth = max(0, int(self.config.hybrid.lexical_candidate_depth))
        return index.search(query_tokens, limit=depth)

    def _lexical_rank(self, query_tokens: list[str], chunks: list[KnowledgeChunk]) -> list[tuple[KnowledgeChunk, float]]:
        """Backward-compatible alias for callers using the old private helper."""
        return self._bm25_rank(query_tokens, chunks)

    def _dense_rank(
        self,
        query: str,
        chunks: list[KnowledgeChunk],
        *,
        collection: str | None = None,
        filters: dict[str, Any] | None = None,
    ) -> list[tuple[KnowledgeChunk, float]]:
        if not query.strip() or not chunks:
            return []
        if self.config.vector_backend == "qdrant":
            if self.embedding_provider is None:
                raise self._backend_error or BackendUnavailableError("BGE-M3 embedding provider is unavailable")
            if self.vector_store is None:
                raise self._backend_error or BackendUnavailableError("Qdrant vector store is unavailable")
            vectors = self.embedding_provider.embed([query])
            if len(vectors) != 1:
                raise BackendUnavailableError("Embedding provider did not return one query vector")
            hits = self.vector_store.search(
                vectors[0],
                limit=max(0, int(self.config.hybrid.dense_candidate_depth)),
                collection=collection,
                filters=filters,
            )
            by_id = {chunk.chunk_id: chunk for chunk in chunks}
            return [
                (by_id[hit.chunk_id], float(hit.score))
                for hit in hits
                if hit.chunk_id in by_id
            ]
        query_vector = _vectorize(query)
        scored = [(chunk, _cosine(query_vector, _vectorize(chunk.text))) for chunk in chunks]
        depth = max(0, int(self.config.hybrid.dense_candidate_depth))
        return [
            (chunk, score)
            for chunk, score in sorted(scored, key=lambda item: (-item[1], item[0].chunk_id))[:depth]
            if score > 0
        ]

    def _rrf(self, ranked_lists: list[list[tuple[KnowledgeChunk, float]]], k: int = 60) -> dict[str, float]:
        scores: dict[str, float] = {}
        for ranked in ranked_lists:
            for rank, (chunk, _) in enumerate(ranked, 1):
                scores[chunk.chunk_id] = scores.get(chunk.chunk_id, 0.0) + 1.0 / (k + rank)
        return scores

    def _expand_with_adjacent_chunks(
        self,
        ranked: list[tuple[KnowledgeChunk, float]],
        chunks: list[KnowledgeChunk],
    ) -> tuple[list[tuple[KnowledgeChunk, float]], dict[str, str]]:
        by_position = {(chunk.doc_id, chunk.ordinal): chunk for chunk in chunks}
        expanded: list[tuple[KnowledgeChunk, float]] = []
        included: set[str] = set()
        adjacent_to: dict[str, str] = {}

        def add(chunk: KnowledgeChunk, score: float, source_chunk_id: str = "") -> None:
            if chunk.chunk_id in included:
                return
            included.add(chunk.chunk_id)
            if source_chunk_id:
                adjacent_to[chunk.chunk_id] = source_chunk_id
            expanded.append((chunk, score))

        for chunk, score in ranked:
            add(chunk, score)
            for ordinal in (chunk.ordinal - 1, chunk.ordinal + 1):
                neighbor = by_position.get((chunk.doc_id, ordinal))
                if neighbor is not None:
                    add(neighbor, score * 0.999, chunk.chunk_id)
        return expanded, adjacent_to

    def search(
        self,
        query: str,
        *,
        collection: str | None = None,
        filters: dict[str, Any] | None = None,
        top_k: int = 8,
        include_scores: bool = True,
    ) -> dict[str, Any]:
        started = time.perf_counter()
        normalized_collection = _safe_component(collection) if collection else None
        chunks = self._candidate_chunks(normalized_collection, filters)
        query_tokens = tokenize(query)
        try:
            bm25 = self._bm25_rank(query_tokens, chunks)
            dense = self._dense_rank(
                query,
                chunks,
                collection=normalized_collection,
                filters=filters,
            )
            rrf_scores = self._rrf(
                [bm25, dense],
                k=max(1, int(self.config.hybrid.rrf_k)),
            )
            by_id = {chunk.chunk_id: chunk for chunk in chunks}
            bm25_scores = {chunk.chunk_id: score for chunk, score in bm25}
            dense_scores = {chunk.chunk_id: score for chunk, score in dense}
            fused = [
                (by_id[chunk_id], score)
                for chunk_id, score in rrf_scores.items()
                if chunk_id in by_id
            ]
            fused.sort(key=lambda item: (-item[1], item[0].chunk_id))

            reranker_scores: dict[str, float] = {}
            reranker_active = False
            if self.reranker_enabled:
                if self.reranker_provider is None:
                    raise self._backend_error or BackendUnavailableError("Configured reranker is unavailable")
                rerank_limit = max(0, int(self.config.reranker.candidate_top_k))
                rerank_candidates = fused[:rerank_limit] if rerank_limit else []
                scores = self.reranker_provider.score(
                    query,
                    [chunk.text for chunk, _ in rerank_candidates],
                )
                if len(scores) != len(rerank_candidates):
                    raise BackendUnavailableError(
                        f"Reranker returned {len(scores)} scores for {len(rerank_candidates)} candidates"
                    )
                reranker_scores = {
                    chunk.chunk_id: float(score)
                    for (chunk, _), score in zip(rerank_candidates, scores)
                }
                reranker_active = bool(rerank_candidates)
                fused.sort(
                    key=lambda item: (
                        -reranker_scores.get(item[0].chunk_id, float("-inf")),
                        -item[1],
                        item[0].chunk_id,
                    )
                )

            fused, adjacent_to = self._expand_with_adjacent_chunks(fused, chunks)
            results = []
            for chunk, fused_score in fused[: max(0, int(top_k))]:
                manifest = self.manifests.get(chunk.doc_id)
                score_payload = {
                    "fusion_score": round(fused_score, 6),
                    "bm25_score": round(bm25_scores.get(chunk.chunk_id, 0.0), 6),
                    "lexical_score": round(bm25_scores.get(chunk.chunk_id, 0.0), 6),
                    "dense_score": round(dense_scores.get(chunk.chunk_id, 0.0), 6),
                }
                if chunk.chunk_id in reranker_scores:
                    score_payload["reranker_score"] = round(reranker_scores[chunk.chunk_id], 6)
                item = {
                    "chunk_id": chunk.chunk_id,
                    "doc_id": chunk.doc_id,
                    "collection": chunk.collection,
                    "citation_id": chunk.chunk_id,
                    "source_ref": chunk.source_ref,
                    "source_uri": manifest.source_uri if manifest else "",
                    "source_url": (manifest.metadata.get("source_url", "") if manifest else ""),
                    "source_type": manifest.source_type if manifest else "",
                    "title": manifest.title if manifest else chunk.source_ref,
                    "heading_path": chunk.heading_path,
                    "snippet": _safe_snippet(chunk.text, query=query),
                    "excerpt": _safe_chunk_excerpt(chunk.text, query=query),
                    "metadata": {
                        "chunking_strategy": chunk.chunking_strategy,
                        "boundary_method": chunk.boundary_method,
                        "start_offset": chunk.start_offset,
                        "end_offset": chunk.end_offset,
                        "adjacent_context": chunk.chunk_id in adjacent_to,
                        "adjacent_to": adjacent_to.get(chunk.chunk_id, ""),
                    },
                }
                if include_scores:
                    item["scores"] = score_payload
                results.append(item)
            return {
                "query": query,
                "collection": normalized_collection,
                "results": results,
                "settings": self._stack_settings(
                    latency_ms=(time.perf_counter() - started) * 1000.0,
                    embedding_active=bool(self.config.vector_backend == "qdrant" and dense),
                    vector_active=bool(self.config.vector_backend == "qdrant" and self.vector_store is not None),
                    reranker_active=reranker_active,
                    candidate_counts={
                        "bm25": len(bm25),
                        "lexical": len(bm25),
                        "dense": len(dense),
                        "fused": len(rrf_scores),
                        "reranked": len(reranker_scores),
                        "returned": len(results),
                    },
                )
                | {
                    "top_k": top_k,
                    "adjacent_chunk_expansion": True,
                },
            }
        except RagBackendError as exc:
            return self._backend_error_result(
                operation="search",
                error=exc,
                query=query,
                collection=normalized_collection,
                results=[],
            )
        except Exception as exc:
            return self._backend_error_result(
                operation="search",
                error=exc,
                query=query,
                collection=normalized_collection,
                results=[],
            )
