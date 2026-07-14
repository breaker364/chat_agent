from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import os
import re
import shlex
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from http import HTTPStatus
from pathlib import Path
from queue import Queue
from threading import Thread
from typing import Any
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from pydantic import BaseModel, Field
from langchain_core.tools import StructuredTool, tool
from langchain_mcp_adapters.sessions import StdioConnection, StreamableHttpConnection
from langchain_mcp_adapters.sessions import create_session
from langchain_mcp_adapters.tools import load_mcp_tools

from langchain_core.messages import HumanMessage, SystemMessage

from .adapters import (
    BingSearchAdapter,
    FetchResult,
    RedirectInfo,
    SearchOptions,
    SearchResult,
    create_fetch_adapter,
    create_search_adapter,
)
from .session_store import SessionStore
from .subagent_runtime import get_subagent_manager
from .subagents import built_in_subagents, run_subagent
from .config import get_runtime_value, load_mcd_mcp_config
from .vision import analyze_image_file
from .feishu_web_login import (
    FeishuWebSessionStore,
    build_feishu_cookies,
    init_feishu_qr_login,
    is_feishu_session_server_valid,
    is_feishu_session_valid,
    poll_feishu_qr_login,
)
from .skills import build_skill_tools

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constants (tunable parameters only — no entity-specific values)
# ---------------------------------------------------------------------------
_ALLOWED_ROOT: str | None = None
_SEARCH_CACHE: dict[str, dict[str, Any]] = {}
_FETCH_CACHE: dict[str, dict[str, Any]] = {}
_SEARCH_CACHE_TTL_SECONDS = 900
_FETCH_CACHE_TTL_SECONDS = 900
_FETCH_TIMEOUT_SECONDS = 15
_MAX_FETCH_REDIRECTS = 10
_PYTHON_RUN_TIMEOUT_SECONDS = 60
_PYTHON_RUN_MAX_TIMEOUT_SECONDS = 120
_PYTHON_OUTPUT_MAX_CHARS = 12_000
_MAX_PARALLEL_SEARCH_ROUTES = 2
_DOWNLOAD_DIR_NAME = str(get_runtime_value("paths", "download_dir", "tmp") or "tmp")
_SUBAGENT_SYNC_TIMEOUT_SECONDS = 120
_SUBAGENT_MANAGER = get_subagent_manager()
_CURRENT_SESSION_ID_ENV = "CHAT_AGENT_SESSION_ID"
_CURRENT_RUN_ID_ENV = "CHAT_AGENT_RUN_ID"
_TOOL_DEDUPE_CACHE: dict[str, dict[str, str]] = {}
_TOOL_DEDUPE_PENDING: dict[str, Any] = {}  # per-run_id -> per-key asyncio.Event + result
_TOOL_DEDUPE_CACHE_MAX_RUNS = 32
_TOOL_CACHEABLE_TTL_SECONDS: dict[str, int] = {
    "list_directory": 300,
    "read_file": 3600,
    "get_file_info": 300,
    "web_search": 900,
    "web_fetch": 900,
    "fetch_webpage": 900,
    "get_subagent_task": 30,
    "feishu_login_status": 60,
    "analyze_image": 3600,
}
_TOOL_SIDE_EFFECT_NAMES = {
    "Agent",
    "write_file",
    "append_file",
    "delete_file",
    "record_script_stage",
    "update_task_plan",
    "record_task_item",
    "update_task_item",
    "record_pitfall",
    "SendMessage",
    "feishu_logout",
}
_TOOL_SIDE_EFFECT_PREFIXES = (
    "write_",
    "append_",
    "delete_",
    "create_",
    "update_",
    "send_",
    "record_",
)
_TOOL_DEFAULT_CACHE_TTL_SECONDS = 0
_TOOL_FAILED_CACHE_TTL_SECONDS = 10


def _run_coro_in_thread(coro: Any) -> Any:
    queue: Queue[tuple[bool, Any]] = Queue(maxsize=1)

    def _target() -> None:
        try:
            result = asyncio.run(coro)
            queue.put((True, result))
        except Exception as exc:
            queue.put((False, exc))

    thread = Thread(target=_target, daemon=True)
    thread.start()
    ok, value = queue.get()
    if ok:
        return value
    raise value


def _run_coro_in_thread_with_timeout(coro: Any, timeout_seconds: int) -> Any:
    queue: Queue[tuple[bool, Any]] = Queue(maxsize=1)

    def _target() -> None:
        try:
            result = asyncio.run(coro)
            queue.put((True, result))
        except Exception as exc:
            queue.put((False, exc))

    thread = Thread(target=_target, daemon=True)
    thread.start()
    try:
        ok, value = queue.get(timeout=max(1, int(timeout_seconds)))
    except Exception as exc:
        raise TimeoutError(f"Operation exceeded {timeout_seconds} seconds.") from exc
    if ok:
        return value
    raise value


def _canonical_json(payload: Any) -> str:
    try:
        return json.dumps(payload or {}, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    except Exception:
        return str(payload)


def _normalize_task_plan_todo(value: Any, index: int) -> dict[str, str]:
    if hasattr(value, "model_dump"):
        value = value.model_dump()
    if isinstance(value, dict):
        item = value
    else:
        raw = str(value or "")
        item = {
            key: match
            for key, match in re.findall(r"([A-Za-z_][A-Za-z0-9_]*)='([^']*)'", raw)
        }
    task_id = str(item.get("task_id") or f"step_{index}").strip()
    content = str(item.get("content") or task_id).strip()
    active_form = str(item.get("activeForm") or item.get("active_form") or content).strip()
    return {
        "task_id": task_id,
        "content": content,
        "activeForm": active_form,
        "status": str(item.get("status") or "pending").strip(),
        "details": str(item.get("details") or "").strip(),
        "result_ref": str(item.get("result_ref") or "").strip(),
    }


def _extract_json_array_from_text(text: str) -> Any | None:
    decoder = json.JSONDecoder()
    for match in re.finditer(r"\[", text or ""):
        try:
            value, _ = decoder.raw_decode(text[match.start():])
        except json.JSONDecodeError:
            continue
        if isinstance(value, list):
            return value
    return None


def _normalize_feishu_batch_write_payload(payload: dict[str, Any]) -> dict[str, Any] | None:
    if str(payload.get("skill_name") or "") != "feishu-personal":
        return None
    request = str(payload.get("request") or "")
    lowered = request.lower()
    if "add-records-batch" not in lowered and "batch" not in lowered and "批量" not in request:
        return None
    records = _extract_json_array_from_text(request)
    if records is None:
        return None
    table_match = re.search(r"\b(tbl[A-Za-z0-9]+)\b", request)
    target_match = re.search(r"https?://[^\s，。；;]+|\b[A-Za-z0-9]{20,}\b", request)
    return {
        "skill_name": "feishu-personal",
        "operation": "bitable.add-records-batch",
        "target": target_match.group(0) if target_match else "",
        "table_id": table_match.group(0) if table_match else "",
        "records_hash": _arguments_hash(records),
    }


def _normalize_tool_payload_for_key(tool_name: str, payload: Any) -> Any:
    if isinstance(payload, dict) and tool_name == "use_skill":
        normalized = _normalize_feishu_batch_write_payload(payload)
        if normalized is not None:
            return normalized
    if tool_name != "update_task_plan" or not isinstance(payload, dict):
        return payload
    todos = payload.get("todos")
    if not isinstance(todos, list):
        todos = []
    return {
        "todos": [
            _normalize_task_plan_todo(todo, index)
            for index, todo in enumerate(todos, 1)
        ],
        "source": str(payload.get("source") or "model").strip(),
        "reason": str(payload.get("reason") or "").strip(),
    }


def _dedupe_key(tool_name: str, payload: Any) -> str:
    return f"{tool_name}:{_canonical_json(_normalize_tool_payload_for_key(tool_name, payload))}"


def _arguments_hash(payload: Any) -> str:
    return "sha256:" + hashlib.sha256(_canonical_json(payload).encode("utf-8")).hexdigest()


def _current_run_id() -> str:
    return os.environ.get(_CURRENT_RUN_ID_ENV, "").strip()


def _current_session_id() -> str:
    return os.environ.get(_CURRENT_SESSION_ID_ENV, "").strip()


def _tool_policy(tool_name: str) -> dict[str, Any]:
    side_effect = tool_name in _TOOL_SIDE_EFFECT_NAMES or any(
        tool_name.startswith(prefix) for prefix in _TOOL_SIDE_EFFECT_PREFIXES
    )
    ttl_seconds = 0 if side_effect else _TOOL_CACHEABLE_TTL_SECONDS.get(tool_name, _TOOL_DEFAULT_CACHE_TTL_SECONDS)
    return {
        "cacheable": ttl_seconds > 0 and not side_effect,
        "ttl_seconds": ttl_seconds,
        "side_effect": side_effect,
    }


def _now_ms() -> int:
    return int(time.time() * 1000)


def _cache_entry_is_reusable(entry: dict[str, Any] | None, policy: dict[str, Any]) -> bool:
    if not entry or not policy.get("cacheable"):
        return False
    if entry.get("status") != "success":
        return False
    expires_at_ms = entry.get("expires_at_ms")
    if isinstance(expires_at_ms, (int, float)) and _now_ms() > int(expires_at_ms):
        return False
    return isinstance(entry.get("content"), str)


def _session_store_for_runtime() -> SessionStore | None:
    session_id = _current_session_id()
    if not session_id:
        return None
    try:
        return SessionStore(_workspace_root())
    except Exception as exc:
        logger.debug("Tool result cache store unavailable: %s", exc)
        return None


def _append_tool_audit(
    *,
    action: str,
    tool_name: str,
    arguments: Any,
    call_key: str,
    dedupe_scope: str = "",
    reused_from_tool_call_id: str = "",
    content: str = "",
    error: str = "",
    latency_ms: int | None = None,
) -> None:
    session_id = _current_session_id()
    if not session_id:
        return
    store = _session_store_for_runtime()
    if store is None:
        return
    payload = {
        "action": action,
        "call_key": call_key,
        "arguments_hash": _arguments_hash(arguments),
        "dedupe_scope": dedupe_scope,
        "reused_from_tool_call_id": reused_from_tool_call_id,
        "latency_ms": latency_ms,
        "error": error,
        "content_preview": content[:500],
    }
    try:
        store.append_tool_event(
            session_id,
            event_type=f"tool_runtime_{action}",
            tool_name=tool_name,
            arguments=arguments,
            content=payload,
        )
    except Exception as exc:
        logger.debug("Failed to append tool audit event: %s", exc)


def _load_conversation_cache_result(tool_name: str, arguments: Any, call_key: str, policy: dict[str, Any]) -> str | None:
    if not policy.get("cacheable"):
        return None
    session_id = _current_session_id()
    store = _session_store_for_runtime()
    if not session_id or store is None:
        return None
    entry = store.get_tool_result_cache_entry(session_id, call_key)
    if not _cache_entry_is_reusable(entry, policy):
        return None
    content = str(entry.get("content") or "")
    _append_tool_audit(
        action="reused",
        tool_name=tool_name,
        arguments=arguments,
        call_key=call_key,
        dedupe_scope="conversation_cache",
        reused_from_tool_call_id=str(entry.get("tool_call_id") or ""),
        content=content,
    )
    return content


def _load_side_effect_repeat_block(tool_name: str, arguments: Any, call_key: str, policy: dict[str, Any]) -> str | None:
    if not policy.get("side_effect"):
        return None
    session_id = _current_session_id()
    store = _session_store_for_runtime()
    if not session_id or store is None:
        return None
    entry = store.get_tool_result_cache_entry(session_id, call_key)
    if not entry or entry.get("status") != "success" or not entry.get("side_effect"):
        return None
    content = _side_effect_repeat_message(tool_name, call_key)
    _append_tool_audit(
        action="blocked",
        tool_name=tool_name,
        arguments=arguments,
        call_key=call_key,
        dedupe_scope="conversation_side_effect",
        reused_from_tool_call_id=str(entry.get("tool_call_id") or ""),
        content=content,
    )
    return content


def _save_conversation_cache_result(
    *,
    tool_name: str,
    arguments: Any,
    call_key: str,
    policy: dict[str, Any],
    content: str,
    status: str,
    error_type: str = "",
) -> None:
    if not policy.get("cacheable") and not policy.get("side_effect") and status == "success":
        return
    session_id = _current_session_id()
    store = _session_store_for_runtime()
    if not session_id or store is None:
        return
    ttl_seconds = int(policy.get("ttl_seconds") or 0)
    if status != "success":
        ttl_seconds = _TOOL_FAILED_CACHE_TTL_SECONDS
    if ttl_seconds <= 0 and not policy.get("side_effect"):
        return
    created_at_ms = _now_ms()
    run_id = _current_run_id()
    entry = {
        "call_key": call_key,
        "tool_name": tool_name,
        "arguments": arguments,
        "content": content,
        "tool_call_id": f"{run_id}:{tool_name}" if run_id else tool_name,
        "created_at_ms": created_at_ms,
        "expires_at_ms": created_at_ms + ttl_seconds * 1000 if ttl_seconds > 0 else None,
        "source_turn_id": run_id,
        "side_effect": bool(policy.get("side_effect")),
        "status": status,
        "error_type": error_type,
    }
    try:
        store.set_tool_result_cache_entry(session_id, call_key, entry)
    except Exception as exc:
        logger.debug("Failed to save tool result cache: %s", exc)


def _side_effect_repeat_message(tool_name: str, call_key: str) -> str:
    return json.dumps(
        {
            "blocked": True,
            "reason": "duplicate_side_effect_tool_call",
            "tool": tool_name,
            "call_key": call_key,
            "message": (
                "Runtime blocked a repeated side-effect tool call with identical arguments. "
                "Ask for explicit confirmation before executing this action again."
            ),
        },
        ensure_ascii=False,
        indent=2,
    )


def _json_objects_from_text(text: str) -> list[dict[str, Any]]:
    objects: list[dict[str, Any]] = []
    raw = (text or "").strip()
    if not raw:
        return objects
    try:
        parsed = json.loads(raw)
        if isinstance(parsed, dict):
            return [parsed]
        if isinstance(parsed, list):
            return [item for item in parsed if isinstance(item, dict)]
    except Exception:
        pass
    decoder = json.JSONDecoder()
    index = 0
    while index < len(raw):
        start = raw.find("{", index)
        if start < 0:
            break
        try:
            parsed, end = decoder.raw_decode(raw[start:])
        except Exception:
            index = start + 1
            continue
        if isinstance(parsed, dict):
            objects.append(parsed)
        index = start + max(end, 1)
    return objects


def _origin_from_text(text: str) -> str:
    match = re.search(r"(https?://[^/\s，。；;]+)", text or "")
    return match.group(1) if match else ""


def _feishu_bitable_url_from_request(request: str, token: str, table_id: str = "") -> str:
    origin = _origin_from_text(request)
    if not origin or not token:
        return ""
    suffix = f"?table={table_id}" if table_id else ""
    return f"{origin}/base/{token}{suffix}"


def _auto_register_tool_result(tool_name: str, arguments: Any, content: str) -> None:
    session_id = _current_session_id()
    if not session_id:
        return
    store = _session_store_for_runtime()
    if store is None:
        return
    args = arguments if isinstance(arguments, dict) else {}
    try:
        for obj in _json_objects_from_text(content):
            if obj.get("success") is False or obj.get("error"):
                continue
            saved_path = obj.get("saved_path") or obj.get("path")
            url = obj.get("url")
            if saved_path or url:
                store.register_artifact(
                    session_id,
                    {
                        "role": "intermediate",
                        "type": "file" if saved_path else "url",
                        "path": str(saved_path or ""),
                        "url": str(url or ""),
                        "summary": str(obj.get("summary") or obj.get("message") or "Tool-produced artifact."),
                        "source_tool": tool_name,
                    },
                )
        if tool_name == "write_file":
            path = str(args.get("path") or "")
            if path:
                store.register_artifact(
                    session_id,
                    {
                        "role": "intermediate",
                        "type": "file",
                        "path": path,
                        "summary": "File written by agent.",
                        "source_tool": tool_name,
                    },
                )
        elif tool_name == "analyze_image":
            path = str(args.get("path") or "")
            stage_name = "image_analysis" + (":" + _arguments_hash(path) if path else "")
            store.record_stage_result(
                session_id,
                stage_name,
                {
                    "status": "completed",
                    "summary": (content or "")[:1000],
                    "result_ref": path,
                    "source_tool": tool_name,
                    "verified": False,
                },
            )
        elif tool_name == "use_skill" and str(args.get("skill_name") or "") == "feishu-personal":
            request = str(args.get("request") or "")
            for obj in _json_objects_from_text(content):
                if obj.get("success") is not True:
                    continue
                if obj.get("records_written") is not None and obj.get("obj_token") and obj.get("table_id"):
                    token = str(obj.get("obj_token") or "")
                    table_id = str(obj.get("table_id") or "")
                    store.record_primary_result(
                        session_id,
                        {
                            "type": "feishu_bitable",
                            "title": "Feishu Bitable records",
                            "status": "written",
                            "url": _feishu_bitable_url_from_request(request, token, table_id),
                            "token": token,
                            "table_id": table_id,
                            "record_count": obj.get("records_written"),
                            "verified": False,
                            "summary": f"Written {obj.get('records_written')} records.",
                            "source_tool": tool_name,
                        },
                    )
                elif obj.get("url") and (obj.get("obj_token") or obj.get("wiki_token")):
                    store.record_primary_result(
                        session_id,
                        {
                            "type": "feishu_bitable",
                            "title": str(obj.get("title") or "Feishu Bitable"),
                            "status": "created",
                            "url": str(obj.get("url") or ""),
                            "token": str(obj.get("obj_token") or ""),
                            "summary": "Feishu Bitable created.",
                            "source_tool": tool_name,
                        },
                    )
                elif obj.get("tableId") and obj.get("total") is not None:
                    progress = store.get_resume_context(session_id)
                    primary = progress.get("primary_result") if isinstance(progress, dict) else None
                    if isinstance(primary, dict) and str(primary.get("table_id") or "") == str(obj.get("tableId")):
                        store.record_primary_result(
                            session_id,
                            {
                                **primary,
                                "status": "verified",
                                "verified": True,
                                "record_count": obj.get("total"),
                                "summary": f"Verified table contains {obj.get('total')} records.",
                                "source_tool": tool_name,
                            },
                        )
    except Exception as exc:
        logger.debug("Auto result registration failed for %s: %s", tool_name, exc)


def _run_cache_for_current_request() -> dict[str, str] | None:
    run_id = _current_run_id()
    if not run_id:
        return None
    if run_id not in _TOOL_DEDUPE_CACHE:
        if len(_TOOL_DEDUPE_CACHE) >= _TOOL_DEDUPE_CACHE_MAX_RUNS:
            oldest = next(iter(_TOOL_DEDUPE_CACHE), None)
            if oldest:
                _TOOL_DEDUPE_CACHE.pop(oldest, None)
                _TOOL_DEDUPE_PENDING.pop(oldest, None)
        _TOOL_DEDUPE_CACHE[run_id] = {}
    return _TOOL_DEDUPE_CACHE[run_id]


def _pending_for_current_request() -> dict[str, asyncio.Event] | None:
    run_id = _current_run_id()
    if not run_id:
        return None
    if run_id not in _TOOL_DEDUPE_PENDING:
        _TOOL_DEDUPE_PENDING[run_id] = {}
    return _TOOL_DEDUPE_PENDING[run_id]


def clear_tool_dedupe_cache(run_id: str | None) -> None:
    if run_id:
        _TOOL_DEDUPE_CACHE.pop(run_id, None)
        _TOOL_DEDUPE_PENDING.pop(run_id, None)


def _wrap_tool_with_run_dedupe(tool_obj: Any) -> Any:
    tool_name = str(getattr(tool_obj, "name", "") or "")
    if not tool_name:
        return tool_obj
    policy = _tool_policy(tool_name)

    def cached_func(**kwargs: Any) -> str:
        cache = _run_cache_for_current_request()
        cache_key = _dedupe_key(tool_name, kwargs)
        if cache is not None and cache_key in cache:
            result_text = cache[cache_key]
            if policy.get("side_effect"):
                result_text = _side_effect_repeat_message(tool_name, cache_key)
                _append_tool_audit(
                    action="blocked",
                    tool_name=tool_name,
                    arguments=kwargs,
                    call_key=cache_key,
                    dedupe_scope="same_run_side_effect",
                    content=result_text,
                )
                return result_text
            _append_tool_audit(
                action="deduped",
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                dedupe_scope="same_run",
                content=result_text,
            )
            return cache[cache_key]

        cached_result = _load_conversation_cache_result(tool_name, kwargs, cache_key, policy)
        if cached_result is not None:
            if cache is not None:
                cache[cache_key] = cached_result
            return cached_result
        blocked_result = _load_side_effect_repeat_block(tool_name, kwargs, cache_key, policy)
        if blocked_result is not None:
            if cache is not None:
                cache[cache_key] = blocked_result
            return blocked_result

        started_at = time.monotonic()
        try:
            result = tool_obj.invoke(kwargs)
            result_text = result if isinstance(result, str) else str(result)
            latency_ms = int((time.monotonic() - started_at) * 1000)
            if cache is not None:
                cache[cache_key] = result_text
            _auto_register_tool_result(tool_name, kwargs, result_text)
            _save_conversation_cache_result(
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                policy=policy,
                content=result_text,
                status="success",
            )
            _append_tool_audit(
                action="executed",
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                content=result_text,
                latency_ms=latency_ms,
            )
            return result_text
        except Exception as exc:
            latency_ms = int((time.monotonic() - started_at) * 1000)
            _save_conversation_cache_result(
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                policy=policy,
                content=str(exc),
                status="failed",
                error_type=type(exc).__name__,
            )
            _append_tool_audit(
                action="failed",
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                error=str(exc),
                latency_ms=latency_ms,
            )
            raise

    async def cached_coroutine(**kwargs: Any) -> str:
        cache = _run_cache_for_current_request()
        cache_key = _dedupe_key(tool_name, kwargs)

        # Fast path: result already cached from a previous call in this run.
        if cache is not None and cache_key in cache:
            result_text = cache[cache_key]
            if policy.get("side_effect"):
                result_text = _side_effect_repeat_message(tool_name, cache_key)
                _append_tool_audit(
                    action="blocked",
                    tool_name=tool_name,
                    arguments=kwargs,
                    call_key=cache_key,
                    dedupe_scope="same_run_side_effect",
                    content=result_text,
                )
                return result_text
            _append_tool_audit(
                action="deduped",
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                dedupe_scope="same_run",
                content=result_text,
            )
            return result_text

        # Pending check: if another task is already executing this exact call,
        # wait for it to complete and reuse its result instead of sending a
        # duplicate request. This prevents parallel duplicates from the same
        # model turn (e.g. calling get-current-date twice simultaneously).
        pending_map = _pending_for_current_request()
        if pending_map is not None and cache_key in pending_map:
            event: asyncio.Event = pending_map[cache_key]
            await event.wait()
            # After the first caller finishes, the result should be in cache.
            if cache is not None and cache_key in cache:
                result_text = cache[cache_key]
                if policy.get("side_effect"):
                    result_text = _side_effect_repeat_message(tool_name, cache_key)
                    _append_tool_audit(
                        action="blocked",
                        tool_name=tool_name,
                        arguments=kwargs,
                        call_key=cache_key,
                        dedupe_scope="same_response_side_effect",
                        content=result_text,
                    )
                    return result_text
                _append_tool_audit(
                    action="deduped",
                    tool_name=tool_name,
                    arguments=kwargs,
                    call_key=cache_key,
                    dedupe_scope="same_response",
                    content=result_text,
                )
                return result_text
            # Fall through and execute if somehow the result wasn't cached.

        cached_result = _load_conversation_cache_result(tool_name, kwargs, cache_key, policy)
        if cached_result is not None:
            if cache is not None:
                cache[cache_key] = cached_result
            return cached_result
        blocked_result = _load_side_effect_repeat_block(tool_name, kwargs, cache_key, policy)
        if blocked_result is not None:
            if cache is not None:
                cache[cache_key] = blocked_result
            return blocked_result

        # Register this call as in-flight so parallel callers can wait.
        if pending_map is not None:
            pending_map[cache_key] = asyncio.Event()

        started_at = time.monotonic()
        try:
            result = await tool_obj.ainvoke(kwargs)
            result_text = result if isinstance(result, str) else str(result)
            if cache is not None:
                cache[cache_key] = result_text
            latency_ms = int((time.monotonic() - started_at) * 1000)
            _auto_register_tool_result(tool_name, kwargs, result_text)
            _save_conversation_cache_result(
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                policy=policy,
                content=result_text,
                status="success",
            )
            _append_tool_audit(
                action="executed",
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                content=result_text,
                latency_ms=latency_ms,
            )
            return result_text
        except Exception as exc:
            latency_ms = int((time.monotonic() - started_at) * 1000)
            _save_conversation_cache_result(
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                policy=policy,
                content=str(exc),
                status="failed",
                error_type=type(exc).__name__,
            )
            _append_tool_audit(
                action="failed",
                tool_name=tool_name,
                arguments=kwargs,
                call_key=cache_key,
                error=str(exc),
                latency_ms=latency_ms,
            )
            raise
        finally:
            if pending_map is not None and cache_key in pending_map:
                pending_map[cache_key].set()
                pending_map.pop(cache_key, None)

    wrapped = StructuredTool.from_function(
        func=cached_func,
        coroutine=cached_coroutine,
        name=tool_name,
        description=getattr(tool_obj, "description", "") or "",
        args_schema=getattr(tool_obj, "args_schema", None),
        return_direct=getattr(tool_obj, "return_direct", False),
    )
    wrapped.metadata = {
        **(getattr(tool_obj, "metadata", None) or {}),
        "dedupe_wrapped": True,
        "dedupe_original_name": tool_name,
    }
    return wrapped

_FETCH_HEADERS = {
    "User-Agent": "Mozilla/5.0",
    "Accept": "text/markdown, text/html, */*;q=0.8",
}

# ---------------------------------------------------------------------------
# File-system helpers (workspace sandboxing)
# ---------------------------------------------------------------------------


def set_allowed_root(path: str | Path) -> None:
    global _ALLOWED_ROOT
    _ALLOWED_ROOT = str(Path(path).resolve())


def _ensure_allowed(path: str) -> Path:
    raw = (path or "").strip() or "."
    candidate = Path(raw)
    root_str = _ALLOWED_ROOT or os.getcwd()
    root = Path(root_str).resolve()
    if candidate.is_absolute():
        target = candidate.resolve()
    else:
        target = (root / candidate).resolve()
    try:
        target.relative_to(root)
    except ValueError as exc:
        raise PermissionError(f"Access denied: {target} is outside {root}") from exc
    return target


def _workspace_root() -> Path:
    return Path(_ALLOWED_ROOT or os.getcwd()).resolve()


def _download_dir() -> Path:
    path = _workspace_root() / _DOWNLOAD_DIR_NAME
    path.mkdir(parents=True, exist_ok=True)
    return path


def _feishu_session_store() -> FeishuWebSessionStore:
    return FeishuWebSessionStore(_workspace_root())


def _current_session_id() -> str:
    return (os.environ.get(_CURRENT_SESSION_ID_ENV) or "").strip()


def _load_feishu_session_payload() -> dict[str, Any]:
    payload = _feishu_session_store().load()
    if not is_feishu_session_valid(payload):
        raise RuntimeError("Feishu web session is missing or expired. Run Feishu login first.")
    return payload or {}


def _parse_json_object(raw: str, field_name: str) -> dict[str, Any]:
    text = (raw or "").strip()
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{field_name} must be valid JSON: {exc}") from exc
    if not isinstance(parsed, dict):
        raise ValueError(f"{field_name} must decode to a JSON object.")
    return parsed


def _validate_feishu_target_url(url: str) -> str:
    normalized = (url or "").strip()
    if not normalized:
        raise ValueError("URL is required.")
    parsed = urlparse(normalized)
    if parsed.scheme != "https":
        raise ValueError("Feishu authenticated requests require an HTTPS URL.")
    hostname = (parsed.hostname or "").lower()
    if not hostname:
        raise ValueError("URL hostname is required.")
    if "feishu" not in hostname and "larksuite" not in hostname:
        raise ValueError("Target URL must be a Feishu/Lark domain.")
    return normalized


def _resolve_python_script_path(path: str) -> Path:
    raw = (path or "").strip()
    if not raw:
        raise ValueError("Python file path is required.")
    candidate = Path(raw)
    if candidate.is_absolute():
        return candidate.resolve()
    return (Path(_workspace_root()) / candidate).resolve()


def _resolve_python_executable(preferred: str = "") -> str:
    configured = preferred.strip()
    candidates = [
        configured,
        (os.environ.get("PYTHON_EXECUTABLE") or "").strip(),
        sys.executable or "",
        shutil.which("python") or "",
        shutil.which("py") or "",
        shutil.which("python3") or "",
    ]
    for candidate in candidates:
        if not candidate:
            continue
        try:
            resolved = Path(candidate).resolve(strict=False)
        except OSError:
            continue
        if resolved.exists():
            return str(resolved)
    raise FileNotFoundError(
        "Python interpreter not found. Set PYTHON_EXECUTABLE or run the agent from a Python environment."
    )


def _kill_process_tree(process: subprocess.Popen[str]) -> None:
    """Terminate the process and any child processes it may have started."""
    try:
        if os.name == "nt":
            subprocess.run(
                ["taskkill", "/F", "/T", "/PID", str(process.pid)],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                check=False,
                timeout=10,
            )
        else:
            os.killpg(process.pid, 9)
    except Exception:
        try:
            process.kill()
        except Exception:
            pass


class RunPythonFileInput(BaseModel):
    """Arguments for running a Python script via the agent tool."""

    path: str = Field(..., description="Path to the Python script. Relative paths use the workspace root.")
    cli_args: str = Field("", description="Command-line arguments passed to the script.")
    timeout_seconds: int = Field(_PYTHON_RUN_TIMEOUT_SECONDS, description="Execution timeout in seconds.")
    working_directory: str = Field("", description="Optional working directory for the script.")
    python_executable: str = Field("", description="Optional Python executable path.")


class AgentToolInput(BaseModel):
    """Arguments for launching a delegated subagent."""

    description: str = Field(..., description="Short task description for the subagent.")
    prompt: str = Field(..., description="Full delegated task prompt.")
    subagent_type: str = Field("general-purpose", description="Subagent type, e.g. general-purpose, Explore, Plan, verification.")
    run_in_background: bool = Field(False, description="Whether to run the subagent in background mode.")


class SendMessageInput(BaseModel):
    """Arguments for sending a message to a background subagent."""

    agent_id: str = Field(..., description="Target background subagent ID.")
    message: str = Field(..., description="Message to deliver to the target subagent.")


class WebFetchInput(BaseModel):
    """Arguments for fetching a webpage or downloading a file."""

    url: str = Field(..., description="Public HTTP or HTTPS URL.")
    prompt: str = Field("", description="Optional extraction prompt for textual webpage content.")
    download: bool = Field(False, description="If true, download the URL as a file into the workspace tmp directory.")


class AnalyzeImageInput(BaseModel):
    """Arguments for analyzing a workspace image with the multimodal model."""

    path: str = Field(..., description="Workspace-relative image path.")
    prompt: str = Field(
        "",
        description="What should be extracted or analyzed from the image.",
    )


class RecentMcdOrdersInput(BaseModel):
    """Arguments for querying recent McDonald's mall orders."""

    last_id: int = Field(0, description="Pagination cursor. Use 0 for the latest page.")
    size: int = Field(10, description="Number of orders to query, capped at 10.")


class ScriptStageInput(BaseModel):
    """Arguments for recording a script processing stage into the current session."""

    stage_name: str = Field(..., description="Short stage name, e.g. fetch, process, write, verify.")
    status: str = Field(..., description="Stage status: planned, running, completed, failed.")
    summary: str = Field("", description="Short human-readable summary of what happened in this stage.")
    artifact_path: str = Field("", description="Optional workspace path to the generated script or output artifact.")


class PrimaryResultInput(BaseModel):
    """Arguments for recording the user-visible primary result of a task."""

    result_type: str = Field(..., description="Result type, e.g. file, feishu_bitable, feishu_doc, report, dataset.")
    title: str = Field("", description="Short result title.")
    url: str = Field("", description="URL where the user can open the result.")
    path: str = Field("", description="Workspace path where the result file is saved.")
    summary: str = Field("", description="Short summary of what the result contains.")
    token: str = Field("", description="Optional service token or id.")
    table_id: str = Field("", description="Optional table id for table-like results.")
    record_count: int | None = Field(None, description="Optional number of records/items in the result.")
    verified: bool = Field(False, description="Whether the result has been verified.")


class StageResultInput(BaseModel):
    """Arguments for recording a reusable completed stage result."""

    stage_name: str = Field(..., description="Stable stage name, e.g. image_analysis or extracted_records.")
    result_ref: str = Field("", description="Path, URL, or token for the reusable stage result.")
    summary: str = Field("", description="Short stage result summary.")
    item_count: int | None = Field(None, description="Optional number of items represented by the stage.")
    verified: bool = Field(False, description="Whether this stage result has been verified.")
    status: str = Field("completed", description="Stage status, usually completed or verified.")


class TaskPlanTodoInput(BaseModel):
    """One item in the complete task plan snapshot."""

    task_id: str = Field(..., description="Stable snake_case id, e.g. fetch_data or verify_output.")
    content: str = Field(..., description="Imperative description of what needs to be done.")
    activeForm: str = Field(..., description="Present-progress wording shown while this task is active.")
    status: str = Field("pending", description="Status: pending, in_progress, completed, failed, or blocked.")
    details: str = Field("", description="Optional concise status details.")
    result_ref: str = Field("", description="Optional result reference such as a file path, URL, tool name, or record id.")


class TaskPlanInput(BaseModel):
    """Arguments for replacing the current session task plan with a full snapshot."""

    todos: list[TaskPlanTodoInput] = Field(..., description="Complete ordered todo list snapshot.")
    source: str = Field("model", description="Plan source label.")
    reason: str = Field(
        "",
        description='Use "plan_changed" only when task IDs must be added, removed, split, or renamed.',
    )


class TaskItemInput(BaseModel):
    """Arguments for saving a task/todo item into the current session."""

    title: str = Field(..., description="Short task title.")
    status: str = Field("pending", description="Task status: pending, running, completed, blocked, or failed.")
    details: str = Field("", description="Concrete progress details, next step, or acceptance criteria.")
    task_id: str = Field("", description="Stable task id. Leave blank to create a new one.")
    artifact_path: str = Field("", description="Optional workspace path related to this task.")


class TaskItemUpdateInput(BaseModel):
    """Arguments for updating an existing session task/todo item."""

    task_id: str = Field(..., description="Task id returned by record_task_item.")
    status: str = Field("", description="Updated status: pending, running, completed, blocked, or failed.")
    title: str = Field("", description="Optional updated task title.")
    details: str = Field("", description="Optional updated progress details.")
    artifact_path: str = Field("", description="Optional workspace path related to this task.")


class PitfallInput(BaseModel):
    """Arguments for saving a pitfall or failed attempt into the current session."""

    summary: str = Field(..., description="Short description of the pitfall, error, or bad assumption.")
    impact: str = Field("", description="What this affected or why it mattered.")
    resolution: str = Field("", description="How it was resolved or what should be tried next.")
    artifact_path: str = Field("", description="Optional workspace path to logs, scripts, or output.")


# ---------------------------------------------------------------------------
# URL / network helpers
# ---------------------------------------------------------------------------


def _validate_public_url(url: str) -> bool:
    try:
        parsed = urlparse((url or "").strip())
    except Exception:
        return False
    if parsed.scheme not in {"http", "https"}:
        return False
    if not parsed.netloc:
        return False
    hostname = (parsed.hostname or "").lower()
    if not hostname:
        return False
    if hostname in {"localhost", "127.0.0.1", "::1"}:
        return False
    if hostname.endswith(".local"):
        return False
    # Reject URLs with username/password
    if parsed.username or parsed.password:
        return False
    # Must have at least one dot in hostname
    if "." not in hostname:
        return False
    return True


def _upgrade_url_if_needed(url: str) -> str:
    parsed = urlparse(url)
    if parsed.scheme == "http":
        return parsed._replace(scheme="https").geturl()
    return url


def _is_permitted_redirect(source_url: str, redirect_url: str) -> bool:
    """Allow same-domain redirects (including www ↔ non-www changes)."""
    try:
        source = urlparse(source_url)
        target = urlparse(redirect_url)
    except Exception:
        return False
    if target.scheme != source.scheme:
        return False
    if target.port != source.port:
        return False
    if target.username or target.password:
        return False
    source_host = re.sub(r"^www\.", "", (source.hostname or "").lower())
    target_host = re.sub(r"^www\.", "", (target.hostname or "").lower())
    return source_host == target_host


def _fetch_url_with_permitted_redirects(
    url: str,
    timeout_seconds: int = _FETCH_TIMEOUT_SECONDS,
) -> requests.Response | dict[str, Any]:
    """Fetch a URL, following same-host redirects, stopping at cross-host ones."""
    current_url = _upgrade_url_if_needed(url)
    for _ in range(_MAX_FETCH_REDIRECTS + 1):
        response = requests.get(
            current_url,
            headers=_FETCH_HEADERS,
            timeout=timeout_seconds,
            allow_redirects=False,
        )
        if response.status_code not in {301, 302, 307, 308}:
            return response
        redirect_location = response.headers.get("Location")
        if not redirect_location:
            break
        redirect_url = urljoin(current_url, redirect_location)
        if _is_permitted_redirect(current_url, redirect_url):
            current_url = redirect_url
            continue
        return {
            "type": "redirect",
            "original_url": current_url,
            "redirect_url": redirect_url,
            "status_code": response.status_code,
        }
    raise requests.TooManyRedirects(f"Too many redirects while fetching {url}")


def _html_to_text(html: str) -> tuple[str, str]:
    """Return (title, body_text) from HTML."""
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))
    return title, text


def _contains_cjk(text: str) -> bool:
    return bool(re.search(r"[\u4e00-\u9fff]", text or ""))


def _contains_latin(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]", text or ""))


def _tokenize_query(query: str) -> list[str]:
    return [token for token in re.split(r"\s+", (query or "").strip()) if token]


def _keyword_fallback_query(query: str) -> str:
    tokens = _tokenize_query(query)
    if not tokens:
        return ""
    return " ".join(tokens[:12])


def _build_cross_language_query(query: str) -> str:
    """Build a keyword-style query in the other language for bilingual search."""
    normalized = (query or "").strip()
    if not normalized:
        return ""

    target_language = "English" if _contains_cjk(normalized) else "Simplified Chinese"
    system_text = (
        "You rewrite web search queries for bilingual retrieval. "
        "Return only a compact search-engine query in the requested language. "
        "Keep product names, model names, years, numbers, and acronyms exact. "
        "Do not add explanations, quotes, or markdown."
    )
    user_text = (
        f"Original query: {normalized}\n"
        f"Target language: {target_language}\n"
        "Rewrite it as a concise search query optimized for web search."
    )
    llm = _get_secondary_llm()
    response = llm.invoke(
        [SystemMessage(content=system_text), HumanMessage(content=user_text)]
    )
    rewritten = response.content if hasattr(response, "content") else str(response)
    candidate = (rewritten or "").strip().strip('"').strip("'")
    if not candidate or candidate.casefold() == normalized.casefold():
        candidate = _keyword_fallback_query(normalized)
        if not candidate or candidate.casefold() == normalized.casefold():
            return ""
    return candidate


def _build_bilingual_queries(query: str) -> list[tuple[str, str]]:
    normalized = (query or "").strip()
    if not normalized:
        return []

    routes: list[tuple[str, str]] = [("primary", normalized)]
    try:
        alternate = _build_cross_language_query(normalized)
    except Exception as exc:
        logger.warning("Cross-language query generation failed: %s", exc)
        alternate = ""
    if alternate and alternate.casefold() != normalized.casefold():
        routes.append(("alternate", alternate))
    return routes[:_MAX_PARALLEL_SEARCH_ROUTES]


def _merge_search_results(
    route_results: list[tuple[str, str, list[SearchResult]]],
    count: int,
) -> tuple[list[SearchResult], list[dict[str, Any]]]:
    """Merge results from multiple search routes, deduplicating by URL.

    Returns at most `count` results, interleaving from different routes
    for diversity (primary first, then alternate, etc.).
    """
    # First pass: collect all unique results per route
    per_route: list[tuple[str, str, list[SearchResult]]] = []
    all_seen: set[str] = set()
    for route_name, route_query, results in route_results:
        unique: list[SearchResult] = []
        for r in results:
            key = (r.url or "").strip().lower()
            if not key or key in all_seen:
                continue
            all_seen.add(key)
            unique.append(r)
        per_route.append((route_name, route_query, unique))

    # Build route_info from all routes
    route_info: list[dict[str, Any]] = [
        {"route": name, "query": q, "result_count": len(items)}
        for name, q, items in per_route
    ]

    # Interleave: round-robin from each route for diversity
    merged: list[SearchResult] = []
    indices = [0] * len(per_route)
    while len(merged) < count:
        added = False
        for i, (_, _, items) in enumerate(per_route):
            if indices[i] < len(items):
                merged.append(items[indices[i]])
                indices[i] += 1
                added = True
                if len(merged) >= count:
                    break
        if not added:
            break

    return merged, route_info


# ---------------------------------------------------------------------------
# File-operation tools
# ---------------------------------------------------------------------------

_READ_FILE_MAX_CHARS = 80_000


def _limit_read_output(text: str) -> str:
    if len(text) <= _READ_FILE_MAX_CHARS:
        return text
    return (
        text[:_READ_FILE_MAX_CHARS]
        + f"\n\n[... truncated {len(text) - _READ_FILE_MAX_CHARS} characters ...]"
    )


def _read_pdf_text(target: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(target))
    return "\n\n".join(
        f"--- Page {index} ---\n{page.extract_text() or ''}"
        for index, page in enumerate(reader.pages, 1)
    )


def _read_docx_text(target: Path) -> str:
    from docx import Document

    document = Document(str(target))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table_index, table in enumerate(document.tables, 1):
        lines.append(f"--- Table {table_index} ---")
        lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(lines)


def _read_workbook_text(target: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(target), read_only=True, data_only=True)
    sections: list[str] = []
    total_chars = 0
    try:
        for sheet in workbook.worksheets:
            header = f"--- Sheet: {sheet.title} ---"
            sections.append(header)
            total_chars += len(header)
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, 500),
                values_only=True,
            ):
                line = "\t".join("" if value is None else str(value) for value in row[:100])
                sections.append(line)
                total_chars += len(line)
                if total_chars >= _READ_FILE_MAX_CHARS:
                    break
            if total_chars >= _READ_FILE_MAX_CHARS:
                break
    finally:
        workbook.close()
    return "\n".join(sections)


def _feishu_crud_script_block_reason(content: str) -> str:
    text = (content or "").lower()
    has_feishu_marker = any(
        marker in text
        for marker in (
            "feishu_web_session.json",
            "/space/api/",
            "/bitable/",
            "lark_tools.commands.bitable",
            "lark_tools.bitable",
        )
    )
    has_write_intent = any(
        marker in text
        for marker in (
            "requests.post",
            "requests.patch",
            "requests.delete",
            "add-record",
            "add_records",
            "set-record",
            "delete-record",
            "cmd_bitable_add",
            "cmd_bitable_set",
            "cmd_bitable_delete",
        )
    )
    if has_feishu_marker and has_write_intent:
        return (
            "Feishu/Lark CRUD scripts are blocked. Use the feishu-personal "
            "skill route instead, for example: "
            "lark bitable add-records-batch <url> <tableId> --json-file <path>."
        )
    return ""


def _blocked_feishu_script_payload(path: str, reason: str) -> str:
    return json.dumps(
        {
            "blocked": True,
            "reason": reason,
            "path": path,
            "suggested_tool": "use_skill",
            "suggested_skill": "feishu-personal",
        },
        ensure_ascii=False,
        indent=2,
    )


@tool
def list_directory(path: str) -> str:
    """List files and folders in a directory. Provide a relative path from the workspace root."""
    if not path or path.strip() == "":
        path = "."
    try:
        target = _ensure_allowed(path)
    except PermissionError as exc:
        return str(exc)
    if not target.is_dir():
        return f"Not a directory: {target}"
    items: list[str] = []
    for entry in sorted(target.iterdir(), key=lambda e: (not e.is_dir(), e.name.lower())):
        suffix = "/" if entry.is_dir() else ""
        items.append(f"{entry.name}{suffix}")
    return "\n".join(items) if items else "(empty directory)"


@tool
def read_file(path: str) -> str:
    """Read text, PDF, DOCX, and XLSX files from the workspace."""
    try:
        target = _ensure_allowed(path)
    except PermissionError as exc:
        return str(exc)
    if not target.is_file():
        return f"File not found: {target}"
    suffix = target.suffix.lower()
    try:
        if suffix == ".pdf":
            return _limit_read_output(_read_pdf_text(target))
        if suffix == ".docx":
            return _limit_read_output(_read_docx_text(target))
        if suffix in {".xlsx", ".xlsm"}:
            return _limit_read_output(_read_workbook_text(target))

        raw = target.read_bytes()
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                return _limit_read_output(raw.decode(encoding))
            except UnicodeDecodeError:
                continue
        return f"[Binary file: {target.name}, size={len(raw)} bytes, path={target}]"
    except Exception as exc:
        return f"Failed to read {target.name}: {exc}"


@tool(args_schema=AnalyzeImageInput)
def analyze_image(path: str, prompt: str = "") -> str:
    """Analyze an image and return visual observations for the main agent."""
    try:
        return analyze_image_file(
            path=path,
            prompt=prompt,
            workspace_root=_workspace_root(),
        )
    except Exception as exc:
        return json.dumps(
            {
                "success": False,
                "error": str(exc),
                "path": path,
            },
            ensure_ascii=False,
        )


@tool
def get_file_info(path: str) -> str:
    """Get metadata about a file or directory (size, modified time, type)."""
    try:
        target = _ensure_allowed(path)
    except PermissionError as exc:
        return str(exc)
    if not target.exists():
        return f"Not found: {target}"
    stat = target.stat()
    kind = "directory" if target.is_dir() else "file"
    return "\n".join(
        [
            f"Name: {target.name}",
            f"Path: {target}",
            f"Type: {kind}",
            f"Size: {stat.st_size:,} bytes",
            f"Modified: {stat.st_mtime}",
        ]
    )


@tool
def write_file(path: str, content: str, overwrite: bool = True) -> str:
    """Create or replace a text file in the workspace. Use overwrite=false to avoid replacing an existing file."""
    try:
        target = _ensure_allowed(path)
    except PermissionError as exc:
        return str(exc)
    if target.suffix.lower() == ".py":
        reason = _feishu_crud_script_block_reason(content or "")
        if reason:
            return _blocked_feishu_script_payload(str(target), reason)
    if target.exists() and target.is_dir():
        return f"Cannot write file because target is a directory: {target}"
    if target.exists() and not overwrite:
        return f"File already exists and overwrite is false: {target}"
    existed_before = target.exists()
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(content or "", encoding="utf-8", newline="\n")
    action = "Updated" if existed_before else "Created"
    return f"{action} file: {target}"


@tool
def append_file(path: str, content: str) -> str:
    """Append text to a file in the workspace. Creates the file if it does not exist."""
    try:
        target = _ensure_allowed(path)
    except PermissionError as exc:
        return str(exc)
    if target.suffix.lower() == ".py":
        existing = target.read_text(encoding="utf-8", errors="ignore") if target.exists() and target.is_file() else ""
        reason = _feishu_crud_script_block_reason(existing + "\n" + (content or ""))
        if reason:
            return _blocked_feishu_script_payload(str(target), reason)
    if target.exists() and target.is_dir():
        return f"Cannot append because target is a directory: {target}"
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("a", encoding="utf-8", newline="\n") as handle:
        handle.write(content or "")
    return f"Appended to file: {target}"


@tool
def delete_file(path: str) -> str:
    """Delete a file in the workspace. Does not delete directories."""
    try:
        target = _ensure_allowed(path)
    except PermissionError as exc:
        return str(exc)
    if not target.exists():
        return f"File not found: {target}"
    if target.is_dir():
        return f"Refusing to delete directory with delete_file: {target}"
    target.unlink()
    return f"Deleted file: {target}"


@tool(args_schema=AgentToolInput)
def Agent(
    description: str,
    prompt: str,
    subagent_type: str = "general-purpose",
    run_in_background: bool = False,
) -> str:
    """Launch a delegated subagent to handle a scoped task."""
    workspace = _workspace_root()
    normalized_type = (subagent_type or "general-purpose").strip()
    if normalized_type not in built_in_subagents():
        normalized_type = "general-purpose"

    if run_in_background:
        session_id = _current_session_id()
        launched = _SUBAGENT_MANAGER.launch(
            prompt=prompt,
            description=description,
            subagent_type=normalized_type,
            workspace_dir=workspace,
            session_id=session_id or None,
        )
        if session_id:
            SessionStore(workspace).add_subagent_task(session_id, launched)
        return json.dumps(
            {
                "status": "async_launched",
                "agent_id": launched["agent_id"],
                "description": description,
                "subagent_type": normalized_type,
            },
            ensure_ascii=False,
            indent=2,
        )

    try:
        payload = _run_coro_in_thread_with_timeout(
            run_subagent(
                prompt=prompt,
                description=description,
                subagent_type=normalized_type,
                workspace_dir=workspace,
            ),
            _SUBAGENT_SYNC_TIMEOUT_SECONDS,
        )
    except TimeoutError as exc:
        session_id = _current_session_id()
        launched = _SUBAGENT_MANAGER.launch(
            prompt=prompt,
            description=description,
            subagent_type=normalized_type,
            workspace_dir=workspace,
            session_id=session_id or None,
        )
        if session_id:
            SessionStore(workspace).add_subagent_task(session_id, launched)
        return json.dumps(
            {
                "status": "timed_out_relaunched",
                "description": description,
                "subagent_type": normalized_type,
                "error": str(exc),
                "replacement_agent_id": launched["agent_id"],
            },
            ensure_ascii=False,
            indent=2,
        )
    session_id = _current_session_id()
    if session_id:
        SessionStore(workspace).add_subagent_task(
            session_id,
            {
                "agent_id": payload["agent_id"],
                "status": "completed",
                "description": description,
                "subagent_type": normalized_type,
                "result": payload["result"],
                "duration_seconds": payload["duration_seconds"],
            },
        )
    return json.dumps(
        {
            "status": "completed",
            "agent_id": payload["agent_id"],
            "description": description,
            "subagent_type": normalized_type,
            "content": payload["result"],
            "duration_seconds": payload["duration_seconds"],
        },
        ensure_ascii=False,
        indent=2,
    )


@tool(args_schema=ScriptStageInput)
def record_script_stage(
    stage_name: str,
    status: str,
    summary: str = "",
    artifact_path: str = "",
) -> str:
    """Record a script processing stage and its result into the current session progress."""
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)
    workspace = _workspace_root()
    payload = {
        "stage_name": (stage_name or "").strip(),
        "status": (status or "").strip(),
        "summary": (summary or "").strip(),
        "artifact_path": (artifact_path or "").strip(),
    }
    SessionStore(workspace).add_script_stage(session_id, payload)
    return json.dumps({"success": True, **payload}, ensure_ascii=False, indent=2)


@tool(args_schema=PrimaryResultInput)
def record_primary_result(
    result_type: str,
    title: str = "",
    url: str = "",
    path: str = "",
    summary: str = "",
    token: str = "",
    table_id: str = "",
    record_count: int | None = None,
    verified: bool = False,
) -> str:
    """Record the primary result location so final answers can return it."""
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)
    payload = {
        "type": (result_type or "artifact").strip(),
        "title": (title or "Result").strip(),
        "url": (url or "").strip(),
        "path": (path or "").strip(),
        "summary": (summary or "").strip(),
        "token": (token or "").strip(),
        "table_id": (table_id or "").strip(),
        "record_count": record_count,
        "verified": bool(verified),
        "status": "verified" if verified else "created",
        "source_tool": "record_primary_result",
    }
    SessionStore(_workspace_root()).record_primary_result(session_id, payload)
    return json.dumps({"success": True, "primary_result": payload}, ensure_ascii=False, indent=2)


@tool(args_schema=StageResultInput)
def record_stage_result(
    stage_name: str,
    result_ref: str = "",
    summary: str = "",
    item_count: int | None = None,
    verified: bool = False,
    status: str = "completed",
) -> str:
    """Record a reusable stage result so continuation does not repeat work."""
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)
    payload = {
        "status": (status or "completed").strip(),
        "result_ref": (result_ref or "").strip(),
        "summary": (summary or "").strip(),
        "item_count": item_count,
        "verified": bool(verified),
        "source_tool": "record_stage_result",
    }
    SessionStore(_workspace_root()).record_stage_result(session_id, stage_name, payload)
    return json.dumps({"success": True, "stage_name": stage_name, "stage_result": payload}, ensure_ascii=False, indent=2)


@tool(args_schema=TaskPlanInput)
def update_task_plan(todos: list[TaskPlanTodoInput], source: str = "model", reason: str = "") -> str:
    """Save or advance the current session task plan with a complete todo snapshot.

    Use this at the start of complex tasks to create the full ordered task list.
    Once a plan exists, the task list is fixed: later calls may only update the
    latest unfinished todo's status/details/result_ref. Completed todos are
    preserved and cannot be downgraded or deleted by later snapshots.
    """
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)

    normalized: list[dict[str, Any]] = []
    in_progress_count = 0
    for index, todo in enumerate(todos, 1):
        item = todo.model_dump() if hasattr(todo, "model_dump") else dict(todo)
        task_id = str(item.get("task_id") or f"step_{index}").strip()
        status = str(item.get("status") or "pending").strip()
        if status == "in_progress":
            in_progress_count += 1
        normalized.append(
            {
                "task_id": task_id,
                "content": str(item.get("content") or task_id).strip(),
                "activeForm": str(item.get("activeForm") or item.get("content") or task_id).strip(),
                "status": status,
                "details": str(item.get("details") or "").strip(),
                "result_ref": str(item.get("result_ref") or "").strip(),
            }
        )

    if len(normalized) >= 2 and in_progress_count > 1:
        return json.dumps(
            {
                "success": False,
                "error": "At most one todo may be in_progress.",
                "in_progress_count": in_progress_count,
            },
            ensure_ascii=False,
            indent=2,
        )

    store = SessionStore(_workspace_root())
    session = store.set_task_plan(
        session_id,
        normalized,
        source=source or "model",
        reason=reason,
    )
    plan = store.load_task_plan(session_id)
    return json.dumps(
        {
            "success": True,
            "task_plan": plan,
            "message": "Task plan snapshot saved. Continue with the pending or in_progress task.",
        },
        ensure_ascii=False,
        indent=2,
    )


@tool(args_schema=TaskItemInput)
def record_task_item(
    title: str,
    status: str = "pending",
    details: str = "",
    task_id: str = "",
    artifact_path: str = "",
) -> str:
    """Save a todo/task item into the current session progress file."""
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)
    payload = {
        "task_id": (task_id or "").strip(),
        "title": (title or "").strip(),
        "status": (status or "pending").strip(),
        "details": (details or "").strip(),
        "artifact_path": (artifact_path or "").strip(),
    }
    session = SessionStore(_workspace_root()).add_task_item(session_id, payload)
    task_items = session.get("task_progress", {}).get("task_items", [])
    saved = next((item for item in reversed(task_items) if item.get("title") == payload["title"]), task_items[-1] if task_items else payload)
    return json.dumps({"success": True, "task": saved}, ensure_ascii=False, indent=2)


@tool(args_schema=TaskItemUpdateInput)
def update_task_item(
    task_id: str,
    status: str = "",
    title: str = "",
    details: str = "",
    artifact_path: str = "",
) -> str:
    """Update a todo/task item saved in the current session progress file."""
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)
    updates = {
        "status": (status or "").strip() or None,
        "title": (title or "").strip() or None,
        "details": (details or "").strip() or None,
        "artifact_path": (artifact_path or "").strip() or None,
    }
    session = SessionStore(_workspace_root()).update_task_item(session_id, task_id, **updates)
    task_items = session.get("task_progress", {}).get("task_items", [])
    saved = next((item for item in task_items if item.get("task_id") == task_id), None)
    return json.dumps({"success": True, "task": saved}, ensure_ascii=False, indent=2)


@tool(args_schema=PitfallInput)
def record_pitfall(
    summary: str,
    impact: str = "",
    resolution: str = "",
    artifact_path: str = "",
) -> str:
    """Save a pitfall, failed attempt, or important lesson into the current session progress file."""
    session_id = _current_session_id()
    if not session_id:
        return json.dumps({"success": False, "error": "No active session id."}, ensure_ascii=False, indent=2)
    payload = {
        "summary": (summary or "").strip(),
        "impact": (impact or "").strip(),
        "resolution": (resolution or "").strip(),
        "artifact_path": (artifact_path or "").strip(),
    }
    SessionStore(_workspace_root()).add_pitfall(session_id, payload)
    return json.dumps({"success": True, **payload}, ensure_ascii=False, indent=2)


@tool
def get_subagent_task(agent_id: str) -> str:
    """Get the status of a background subagent task."""
    task = _SUBAGENT_MANAGER.get_task((agent_id or "").strip(), workspace_dir=_workspace_root())
    if task is None:
        return json.dumps({"error": "Subagent task not found."}, ensure_ascii=False, indent=2)
    return json.dumps(task, ensure_ascii=False, indent=2)


@tool(args_schema=SendMessageInput)
def SendMessage(agent_id: str, message: str) -> str:
    """Send a message to a background subagent mailbox."""
    ok = _SUBAGENT_MANAGER.send_message(
        (agent_id or "").strip(),
        (message or "").strip(),
        workspace_dir=_workspace_root(),
    )
    if not ok:
        return json.dumps({"success": False, "error": "Target subagent not found."}, ensure_ascii=False, indent=2)
    return json.dumps(
        {
            "success": True,
            "agent_id": agent_id,
            "message": "Message queued for subagent.",
        },
        ensure_ascii=False,
        indent=2,
    )


@tool(args_schema=RunPythonFileInput)
def run_python_file(
    path: str,
    cli_args: str = "",
    timeout_seconds: int = _PYTHON_RUN_TIMEOUT_SECONDS,
    working_directory: str = "",
    python_executable: str = "",
) -> str:
    """Run a Python file for local verification. Relative paths use the workspace root; absolute paths may be outside it."""
    try:
        target = _resolve_python_script_path(path)
    except Exception as exc:
        return json.dumps({"path": path, "error": str(exc)}, ensure_ascii=False, indent=2)

    if not target.is_file():
        return json.dumps({"path": str(target), "error": "File not found."}, ensure_ascii=False, indent=2)
    if target.suffix.lower() != ".py":
        return json.dumps({"path": str(target), "error": "Only .py files can be executed."}, ensure_ascii=False, indent=2)
    try:
        script_text = target.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        script_text = ""
    reason = _feishu_crud_script_block_reason(script_text)
    if reason:
        return _blocked_feishu_script_payload(str(target), reason)

    try:
        timeout_value = min(_PYTHON_RUN_MAX_TIMEOUT_SECONDS, max(1, int(timeout_seconds)))
    except (TypeError, ValueError):
        timeout_value = _PYTHON_RUN_TIMEOUT_SECONDS

    try:
        resolved_python_executable = _resolve_python_executable(python_executable)
    except Exception as exc:
        return json.dumps(
            {"path": str(target), "error": str(exc)},
            ensure_ascii=False,
            indent=2,
        )

    if working_directory.strip():
        try:
            cwd_path = Path(working_directory).expanduser()
            if not cwd_path.is_absolute():
                cwd_path = (_workspace_root() / cwd_path).resolve()
            else:
                cwd_path = cwd_path.resolve()
        except Exception as exc:
            return json.dumps(
                {"path": str(target), "error": f"Invalid working_directory: {exc}"},
                ensure_ascii=False,
                indent=2,
            )
    else:
        cwd_path = target.parent if target.parent.exists() else _workspace_root()

    if not cwd_path.exists() or not cwd_path.is_dir():
        return json.dumps(
            {
                "path": str(target),
                "error": f"Working directory not found or not a directory: {cwd_path}",
            },
            ensure_ascii=False,
            indent=2,
        )

    command = [resolved_python_executable, str(target), *shlex.split(cli_args or "", posix=False)]
    started_at = time.monotonic()
    creationflags = getattr(subprocess, "CREATE_NEW_PROCESS_GROUP", 0) if os.name == "nt" else 0
    popen_kwargs: dict[str, Any] = {
        "cwd": str(cwd_path),
        "stdout": subprocess.PIPE,
        "stderr": subprocess.PIPE,
        "stdin": subprocess.DEVNULL,
        "text": True,
        "encoding": "utf-8",
        "errors": "replace",
        "shell": False,
    }
    if os.name == "nt":
        popen_kwargs["creationflags"] = creationflags
    else:
        popen_kwargs["start_new_session"] = True

    try:
        process = subprocess.Popen(command, **popen_kwargs)
    except Exception as exc:
        return json.dumps(
            {"path": str(target), "command": command, "error": str(exc)},
            ensure_ascii=False,
            indent=2,
        )

    try:
        stdout_text, stderr_text = process.communicate(timeout=timeout_value)
    except subprocess.TimeoutExpired:
        _kill_process_tree(process)
        try:
            process.wait(timeout=2)
        except Exception:
            pass
        return json.dumps(
            {
                "path": str(target),
                "command": command,
                "cwd": str(cwd_path),
                "timed_out": True,
                "timeout_seconds": timeout_value,
                "duration_seconds": round(time.monotonic() - started_at, 3),
                "stdout": "",
                "stderr": "",
            },
            ensure_ascii=False,
            indent=2,
        )
    except Exception as exc:
        return json.dumps(
            {"path": str(target), "command": command, "error": str(exc)},
            ensure_ascii=False,
            indent=2,
        )

    return json.dumps(
        {
            "path": str(target),
            "command": command,
            "cwd": str(cwd_path),
            "exit_code": process.returncode,
            "duration_seconds": round(time.monotonic() - started_at, 3),
            "stdout": (stdout_text or "")[:_PYTHON_OUTPUT_MAX_CHARS],
            "stderr": (stderr_text or "")[:_PYTHON_OUTPUT_MAX_CHARS],
        },
        ensure_ascii=False,
        indent=2,
    )


@tool
def feishu_login_status() -> str:
    """Return whether a reusable Feishu web session is stored locally and actually valid on the server.

    This performs a real HTTP probe against the Feishu server to verify the
    session cookie is still accepted. A purely local timestamp check would
    give false positives for server-side expired/revoked sessions.
    """
    payload = _feishu_session_store().load()
    local_valid = is_feishu_session_valid(payload)
    has_session = bool(payload and payload.get("session"))
    server = {"checked": False, "valid": False, "reason": "skipped"}
    if has_session:
        try:
            server = is_feishu_session_server_valid(payload)
        except Exception as exc:
            server = {"checked": True, "valid": False, "reason": f"probe_exception: {exc}"}
    logged_in = has_session and local_valid and server.get("valid", False)
    result = {
        "logged_in": logged_in,
        "has_session": has_session,
        "issued_at": payload.get("issued_at") if payload else None,
        "metadata": payload.get("metadata", {}) if payload else {},
        "local_check": {"valid": local_valid, "reason": "timestamp_check" if local_valid else "expired_or_missing"},
        "server_check": server,
    }
    if not logged_in and server.get("valid") is False and server.get("reason", "").startswith("server_"):
        result["action_required"] = (
            "The Feishu session is invalid on the server side (e.g. CSRF token was rotated, session was revoked, "
            "or the cookie expired server-side). Run Feishu QR login again to obtain a fresh session."
        )
    return json.dumps(result, ensure_ascii=False, indent=2)


@tool
def feishu_logout() -> str:
    """Clear the locally stored Feishu web session."""
    _feishu_session_store().clear()
    return json.dumps({"success": True, "message": "Feishu web session cleared."}, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Web search tool (adapter-based, like Claude Code)
# ---------------------------------------------------------------------------


def _format_search_results(query: str, results: list[SearchResult], adapter_name: str, duration_ms: float) -> str:
    """Format search results as structured JSON — compatible with Claude Code output style."""
    output: dict[str, Any] = {
        "query": query,
        "engine": adapter_name,
        "duration_seconds": round(duration_ms / 1000, 3),
        "result_count": len(results),
        "results": [
            {
                "title": r.title,
                "url": r.url,
                "snippet": r.snippet,
            }
            for r in results
        ],
    }
    if not results:
        output["note"] = "No search results found. Try a different query or narrower terms."
    return json.dumps(output, ensure_ascii=False, indent=2)


@tool
def web_search(
    query: str,
    count: int = 8,
    allowed_domains: list[str] | None = None,
    blocked_domains: list[str] | None = None,
) -> str:
    """Search the web and return structured results (titles, URLs, snippets).

    The search adapter is selected automatically — Tavily if configured via
    TAVILY_BASE_URL, otherwise Bing HTML scraping (no API key required).

    Use `allowed_domains` to restrict results to specific domains.
    Use `blocked_domains` to exclude specific domains.
    """
    # Cache check
    cache_key = json.dumps(
        {
            "query": query.strip(),
            "count": count,
            "allowed_domains": allowed_domains or [],
            "blocked_domains": blocked_domains or [],
        },
        ensure_ascii=False,
        sort_keys=True,
    )
    now = time.time()
    cached = _SEARCH_CACHE.get(cache_key)
    if cached and now - cached["ts"] < _SEARCH_CACHE_TTL_SECONDS:
        return json.dumps(cached["payload"], ensure_ascii=False, indent=2)

    start = time.monotonic()
    routes = _build_bilingual_queries(query.strip())
    options = SearchOptions(
        allowed_domains=allowed_domains,
        blocked_domains=blocked_domains,
        num_results=count,
    )

    try:
        adapter = create_search_adapter()
        route_results: list[tuple[str, str, list[SearchResult]]] = []
        with ThreadPoolExecutor(max_workers=max(1, len(routes))) as executor:
            future_map = {
                executor.submit(adapter.search, route_query, options): (route_name, route_query)
                for route_name, route_query in routes
            }
            for future in as_completed(future_map):
                route_name, route_query = future_map[future]
                route_results.append((route_name, route_query, future.result()))
        route_results.sort(key=lambda item: 0 if item[0] == "primary" else 1)
        results, route_info = _merge_search_results(route_results, count)
        duration_ms = (time.monotonic() - start) * 1000
        payload_str = _format_search_results(query, results, adapter.name, duration_ms)
        payload = json.loads(payload_str)
        payload["search_routes"] = route_info
        payload_str = json.dumps(payload, ensure_ascii=False, indent=2)
    except Exception as exc:
        requested_engine = getattr(adapter, "name", "?")
        logger.warning("Search adapter '%s' failed: %s", requested_engine, exc)
        # Fallback to Bing if Tavily fails
        adapter = BingSearchAdapter()
        try:
            route_results = []
            with ThreadPoolExecutor(max_workers=max(1, len(routes))) as executor:
                future_map = {
                    executor.submit(adapter.search, route_query, options): (route_name, route_query)
                    for route_name, route_query in routes
                }
                for future in as_completed(future_map):
                    route_name, route_query = future_map[future]
                    route_results.append((route_name, route_query, future.result()))
            route_results.sort(key=lambda item: 0 if item[0] == "primary" else 1)
            results, route_info = _merge_search_results(route_results, count)
            duration_ms = (time.monotonic() - start) * 1000
            payload_str = _format_search_results(query, results, adapter.name, duration_ms)
            payload = json.loads(payload_str)
            payload["requested_engine"] = requested_engine
            payload["fallback_reason"] = str(exc)
            payload["search_routes"] = route_info
            payload_str = json.dumps(payload, ensure_ascii=False, indent=2)
        except Exception as fallback_exc:
            duration_ms = (time.monotonic() - start) * 1000
            payload_str = json.dumps(
                {
                    "query": query,
                    "engine": "none",
                    "duration_seconds": round(duration_ms / 1000, 3),
                    "result_count": 0,
                    "results": [],
                    "error": f"Search failed: {fallback_exc}",
                },
                ensure_ascii=False,
                indent=2,
            )

    payload = json.loads(payload_str)
    _SEARCH_CACHE[cache_key] = {"ts": now, "payload": payload}
    return payload_str


# ---------------------------------------------------------------------------
# Web fetch tool (adapter-based, like Claude Code)
# ---------------------------------------------------------------------------

# Cached secondary LLM for applying prompt to fetched content
_SECONDARY_LLM: Any = None
_FETCH_CONTENT_MAX_CHARS = 10_000  # max chars fed to secondary model


def _get_secondary_llm():
    """Return a cached ChatDeepSeek instance for content processing."""
    global _SECONDARY_LLM
    if _SECONDARY_LLM is None:
        from .config import create_chat_deepseek, load_llm_config

        cfg = load_llm_config()
        _SECONDARY_LLM = create_chat_deepseek(cfg, temperature=0.0, streaming=False)
    return _SECONDARY_LLM


def _apply_prompt_to_content(prompt: str, content: str) -> str:
    """Process fetched content with the user's prompt via a secondary LLM call.

    Mirrors Claude Code's applyPromptToMarkdown / makeSecondaryModelPrompt pattern.
    """
    truncated = content[:_FETCH_CONTENT_MAX_CHARS]
    if len(content) > _FETCH_CONTENT_MAX_CHARS:
        truncated += "\n\n[Content truncated due to length...]"

    system_text = (
        "You are a precise content extraction assistant. "
        "Answer ONLY based on the provided web page content. "
        "If the content does not contain the requested information, say so clearly. "
        "Keep quotes under 125 characters. Do not reproduce song lyrics."
    )
    user_text = (
        f"Web page content:\n---\n{truncated}\n---\n\n"
        f"{prompt}\n\n"
        f"Provide a concise response based only on the content above. "
        f"Include specific details, numbers, and names where available."
    )

    llm = _get_secondary_llm()
    response = llm.invoke(
        [SystemMessage(content=system_text), HumanMessage(content=user_text)]
    )
    return response.content if hasattr(response, "content") else str(response)


def _format_fetch_success(fetch_result: FetchResult, prompt: str) -> str:
    """Format a successful fetch as JSON — matches Claude Code output schema."""
    content = fetch_result.content
    if prompt and prompt.strip():
        try:
            result = _apply_prompt_to_content(prompt, content)
        except Exception as exc:
            logger.warning("Secondary model call failed: %s — returning raw content", exc)
            result = content[:_FETCH_CONTENT_MAX_CHARS]
    else:
        result = content[:_FETCH_CONTENT_MAX_CHARS]

    payload: dict[str, Any] = {
        "url": fetch_result.url,
        "code": fetch_result.code,
        "code_text": fetch_result.code_text,
        "bytes": fetch_result.bytes,
        "result": result,
        "content_type": fetch_result.content_type,
        "duration_seconds": round(fetch_result.duration_ms / 1000, 3),
        "prompt": prompt,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _format_fetch_redirect(redirect: RedirectInfo, prompt: str) -> str:
    """Format a cross-host redirect notice as JSON."""
    payload = {
        "url": redirect.original_url,
        "code": redirect.status_code,
        "code_text": HTTPStatus(redirect.status_code).phrase,
        "result": (
            "REDIRECT DETECTED: The URL redirects to a different host.\n"
            f"Original URL: {redirect.original_url}\n"
            f"Redirect URL: {redirect.redirect_url}\n"
            "Please call web_fetch again with the redirect URL if you want that page."
        ),
        "duration_seconds": 0,
        "prompt": prompt,
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _infer_download_filename(url: str, content_type: str | None = None) -> str:
    parsed = urlparse(url)
    name = Path(parsed.path).name.strip()
    extension = ""
    if content_type:
        lowered = content_type.lower()
        if "pdf" in lowered:
            extension = ".pdf"
        elif "zip" in lowered:
            extension = ".zip"
        elif "json" in lowered:
            extension = ".json"
        elif "csv" in lowered:
            extension = ".csv"
        elif "png" in lowered:
            extension = ".png"
        elif "jpeg" in lowered or "jpg" in lowered:
            extension = ".jpg"
    if name:
        if "." in name or not extension:
            return name
        return f"{name}{extension}"
    return f"download-{int(time.time() * 1000)}{extension}"


def _download_binary_url(url: str) -> str:
    started_at = time.monotonic()
    response = _fetch_url_with_permitted_redirects(url)
    if isinstance(response, dict):
        return _format_fetch_redirect(
            RedirectInfo(
                original_url=response["original_url"],
                redirect_url=response["redirect_url"],
                status_code=response["status_code"],
            ),
            prompt="",
        )

    response.raise_for_status()
    content_type = response.headers.get("Content-Type", "")
    filename = _infer_download_filename(str(response.url), content_type)
    target = _download_dir() / filename
    target.write_bytes(response.content)
    payload = {
        "url": str(response.url),
        "code": response.status_code,
        "code_text": response.reason,
        "bytes": len(response.content),
        "content_type": content_type or "application/octet-stream",
        "downloaded": True,
        "saved_path": str(target),
        "duration_seconds": round(time.monotonic() - started_at, 3),
    }
    return json.dumps(payload, ensure_ascii=False, indent=2)


def _is_safe_tool_schema(tool: Any) -> bool:
    try:
        schema = getattr(tool, "args_schema", None)
        if isinstance(schema, dict):
            payload = schema
        elif hasattr(schema, "model_json_schema"):
            payload = schema.model_json_schema()
        elif hasattr(schema, "schema"):
            payload = schema.schema()
        else:
            payload = {}
        return payload.get("type") in {None, "object"}
    except Exception:
        return False


@tool(args_schema=WebFetchInput)
def web_fetch(url: str, prompt: str = "", download: bool = False) -> str:
    """Fetch a webpage, process with the prompt, and return a result.

    - Fetches the URL, converts HTML to text/Markdown.
    - If `download=true`, downloads the URL as a file into the workspace `tmp/` directory.
    - If Tavily Extract is configured (TAVILY_BASE_URL), returns clean Markdown directly.
    - When `prompt` is provided, a secondary model processes the content
      to answer your prompt, and the answer is returned in the `result` field.
    - When `prompt` is empty, `result` contains the raw content (truncated).
    - HTTP URLs are automatically upgraded to HTTPS.
    - Cross-host redirects are reported (not followed).
    - Results are cached for 15 minutes.
    """
    if not _validate_public_url(url):
        return json.dumps({"url": url, "error": "Invalid URL."}, ensure_ascii=False, indent=2)

    if download:
        try:
            return _download_binary_url(url)
        except requests.Timeout:
            return json.dumps(
                {"url": url, "error": f"Timed out after {_FETCH_TIMEOUT_SECONDS}s while downloading the file."},
                ensure_ascii=False,
                indent=2,
            )
        except Exception as exc:
            return json.dumps({"url": url, "error": str(exc)}, ensure_ascii=False, indent=2)

    # Cache check (keyed by url only — prompt variations share the same fetched content)
    cache_key = json.dumps({"url": url}, ensure_ascii=False, sort_keys=True)
    now = time.time()
    cached = _FETCH_CACHE.get(cache_key)
    cached_content: str | None = None
    if cached and now - cached["ts"] < _FETCH_CACHE_TTL_SECONDS:
        cached_content = cached.get("content")

    started_at = time.monotonic()

    if cached_content is not None:
        # Re-apply prompt to cached content
        fetch_result = FetchResult(
            url=url,
            content=cached_content,
            content_type=cached.get("content_type", "text/plain"),
            bytes=cached.get("bytes", len(cached_content.encode("utf-8"))),
            code=cached.get("code", 200),
            code_text=cached.get("code_text", "OK"),
            duration_ms=0,
        )
        return _format_fetch_success(fetch_result, prompt)

    # Try Tavily Extract first if available
    fetch_adapter = create_fetch_adapter()
    if fetch_adapter is not None:
        try:
            result = fetch_adapter.fetch(url)
            if isinstance(result, RedirectInfo):
                return _format_fetch_redirect(result, prompt)
            # Cache the raw content
            _FETCH_CACHE[cache_key] = {
                "ts": now,
                "content": result.content,
                "content_type": result.content_type,
                "bytes": result.bytes,
                "code": result.code,
                "code_text": result.code_text,
            }
            return _format_fetch_success(result, prompt)
        except Exception as exc:
            logger.warning("Tavily fetch failed for %s: %s — falling back to direct HTTP", url, exc)

    # Direct HTTP fetch
    try:
        response = _fetch_url_with_permitted_redirects(url)
    except requests.Timeout:
        return json.dumps(
            {"url": url, "error": f"Timed out after {_FETCH_TIMEOUT_SECONDS}s while fetching the page."},
            ensure_ascii=False,
            indent=2,
        )
    except Exception as exc:
        return json.dumps({"url": url, "error": str(exc)}, ensure_ascii=False, indent=2)

    if isinstance(response, dict):
        return _format_fetch_redirect(
            RedirectInfo(
                original_url=response["original_url"],
                redirect_url=response["redirect_url"],
                status_code=response["status_code"],
            ),
            prompt,
        )

    try:
        response.raise_for_status()
        response.encoding = response.apparent_encoding or response.encoding
        _, text = _html_to_text(response.text)
        duration_ms = (time.monotonic() - started_at) * 1000

        fetch_result = FetchResult(
            url=str(response.url),
            content=text,
            content_type="text/plain",
            bytes=len(response.content),
            code=response.status_code,
            code_text=response.reason,
            duration_ms=duration_ms,
        )
        # Cache the raw content (without prompt-specific result)
        _FETCH_CACHE[cache_key] = {
            "ts": now,
            "content": text,
            "content_type": "text/plain",
            "bytes": len(response.content),
            "code": response.status_code,
            "code_text": response.reason,
        }
        return _format_fetch_success(fetch_result, prompt)
    except Exception as exc:
        return json.dumps(
            {"url": url, "error": f"HTTP {response.status_code}: {exc}"},
            ensure_ascii=False,
            indent=2,
        )


fetch_webpage = web_fetch


# ---------------------------------------------------------------------------
# 12306 MCP tools (unchanged)
# ---------------------------------------------------------------------------


async def load_12306_tools(
    env: dict[str, str] | None = None,
    cwd: str | None = None,
) -> list[Any]:
    """Start the 12306-mcp server and return the LangChain tools."""
    connection = StdioConnection(
        transport="stdio",
        command="npx",
        args=["-y", "12306-mcp"],
        cwd=cwd or os.getcwd(),
        encoding="utf-8",
        encoding_error_handler="replace",
        env=env or dict(os.environ),
    )
    return await load_mcp_tools(session=None, connection=connection)


async def load_mcd_tools(
    config_path: str | Path | None = None,
) -> list[Any]:
    """Connect to the McDonald's MCP server via Streamable HTTP."""
    cfg = load_mcd_mcp_config(config_path)
    token = (cfg.get("token") or "").strip()
    if not token:
        raise ValueError("MCD_MCP_TOKEN is not configured.")
    connection: StreamableHttpConnection = {
        "transport": "streamable_http",
        "url": str(cfg["url"]),
        "headers": {
            "Authorization": f"Bearer {token}",
        },
    }
    tools = await load_mcp_tools(
        session=None,
        connection=connection,
        server_name="mcd-mcp",
    )
    safe_tools: list[Any] = []
    for tool in tools:
        if getattr(tool, "name", "") == "mall-order-list":
            logger.warning("Replacing mcd MCP tool with local wrapper due to schema incompatibility: %s", tool.name)
            continue
        if _is_safe_tool_schema(tool):
            safe_tools.append(tool)
        else:
            logger.warning("Skipping incompatible mcd MCP tool schema: %s", getattr(tool, "name", "<unknown>"))
    return safe_tools


async def _query_recent_mcd_orders_direct(last_id: int = 0, size: int = 10) -> dict[str, Any]:
    cfg = load_mcd_mcp_config()
    token = (cfg.get("token") or "").strip()
    if not token:
        raise ValueError("MCD_MCP_TOKEN is not configured.")
    connection: StreamableHttpConnection = {
        "transport": "streamable_http",
        "url": str(cfg["url"]),
        "headers": {
            "Authorization": f"Bearer {token}",
        },
    }
    async with create_session(connection) as session:
        from mcp import types as mcp_types

        await session.initialize()
        result = await session.send_request(
            mcp_types.ClientRequest(
                mcp_types.CallToolRequest(
                    params=mcp_types.CallToolRequestParams(
                        name="mall-order-list",
                        arguments={
                            "lastId": last_id,
                            "size": min(max(int(size), 1), 10),
                        },
                    )
                )
            ),
            mcp_types.CallToolResult,
        )
        payload = result.model_dump() if hasattr(result, "model_dump") else result.dict()
        return payload


@tool(args_schema=RecentMcdOrdersInput)
def query_recent_mcd_orders(last_id: int = 0, size: int = 10) -> str:
    """Query the user's recent McDonald's mall orders."""
    payload = _run_coro_in_thread(_query_recent_mcd_orders_direct(last_id=last_id, size=size))
    return json.dumps(payload, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Tool collections
# ---------------------------------------------------------------------------

_FILE_TOOLS: list[Any] = [
    list_directory,
    read_file,
    analyze_image,
    get_file_info,
    write_file,
    append_file,
    delete_file,
    run_python_file,
]
_SEARCH_TOOLS: list[Any] = [web_search, web_fetch]
_AGENT_TOOLS: list[Any] = [
    Agent,
    get_subagent_task,
    SendMessage,
    record_script_stage,
    record_primary_result,
    record_stage_result,
    update_task_plan,
    record_task_item,
    update_task_item,
    record_pitfall,
    query_recent_mcd_orders,
    feishu_login_status,
    feishu_logout,
]


async def get_all_tools(
    workspace_dir: str | Path | None = None,
) -> list[Any]:
    """Return the complete tool list: local search, file ops, and 12306 tools."""
    if workspace_dir is not None:
        set_allowed_root(workspace_dir)
    workspace = Path(workspace_dir or os.getcwd()).resolve()
    tools = list(_FILE_TOOLS) + list(_SEARCH_TOOLS) + list(_AGENT_TOOLS)
    tools.extend(build_skill_tools(workspace))
    try:
        ticket_tools = await load_12306_tools(cwd=str(workspace_dir or os.getcwd()))
        tools.extend(ticket_tools)
    except Exception as exc:
        logger.warning("12306 MCP tools unavailable: %s", exc)
    try:
        mcd_tools = await load_mcd_tools()
        tools.extend(mcd_tools)
    except Exception as exc:
        logger.warning("mcd MCP tools unavailable: %s", exc)
    return [_wrap_tool_with_run_dedupe(item) for item in tools]
